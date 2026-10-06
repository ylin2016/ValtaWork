"""Nights for sale priced below the listing's Wheelhouse minimum ->
output/min_price_check_2026-10-06.xlsx (summary by listing + night detail)."""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
m = pd.read_pickle(ROOT / "data" / "min_price_check.pkl")
m["month"] = m.date.str[:7]
FILL = {"Systemic": "E53935", "Seasonal (Jan–Mar)": "FDD835", "Partial": "FB8C00", "Last-minute only": "B0BEC5"}


def pattern(g):
    b = g[g.below]
    share = len(b) / len(g)
    if b.empty:
        return None
    if b.days_out.max() <= 30:
        return "Last-minute only"
    early = g[g.date < "2027-01-01"]
    if len(early) and early.below.mean() < 0.05 and b.date.min() >= "2027-01-01":
        return "Seasonal (Jan–Mar)"
    return "Systemic" if share >= 0.75 else "Partial"


rows = []
for (lid, name), g in m.groupby(["listing_id", "nickname"]):
    p = pattern(g)
    if not p:
        continue
    b = g[g.below]
    by_m = (g.groupby("month").below.mean() * 100).round(0)
    rows.append({"listing": name, "pattern": p, "wh_min_price": g.min_price.iloc[0],
                 "wh_adjustment": g.base_price_adjustment.iloc[0], "nights_for_sale": len(g),
                 "nights_below_min": len(b), "pct_below": round(len(b) / len(g), 3),
                 "lowest_price": b.price.min(), "median_price_below": b.price.median(),
                 "total_gap_to_min": round(b.gap.sum()), "first_night_below": b.date.min(),
                 "last_night_below": b.date.max(), "median_days_out_below": b.days_out.median(),
                 "guesty_price_matches_wh_rec": round(g.match.mean(), 3),
                 **{f"% below {k}": v for k, v in by_m.items()}})
s = pd.DataFrame(rows)
order = {"Systemic": 0, "Partial": 1, "Seasonal (Jan–Mar)": 2, "Last-minute only": 3}
s = s.sort_values(["pattern", "total_gap_to_min"], key=lambda c: c.map(order) if c.name == "pattern" else -c)
d = m[m.below][["nickname", "date", "days_out", "price", "rec", "min_price", "gap", "base_price_adjustment"]]
d = d.rename(columns={"nickname": "listing", "price": "guesty_price", "rec": "wh_recommended",
                      "gap": "gap_to_min"}).sort_values(["listing", "date"])
out = ROOT / "output" / "min_price_check_2026-10-06.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    s.to_excel(xw, sheet_name="by_listing", index=False)
    d.to_excel(xw, sheet_name="nights_below_min", index=False)
    for ws in xw.book.worksheets:
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 60)), 1):
            ws.column_dimensions[get_column_letter(i)].width = min(28, max(9, max(len(str(x.value or "")) for x in col) + 2))
    ws = xw.book["by_listing"]
    hdr = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        cell = row[hdr.index("pattern")]
        cell.fill = PatternFill("solid", fgColor=FILL[cell.value])
        row[hdr.index("pct_below")].number_format = "0%"
        row[hdr.index("guesty_price_matches_wh_rec")].number_format = "0%"
        for k in ("wh_min_price", "lowest_price", "median_price_below", "total_gap_to_min"):
            row[hdr.index(k)].number_format = '"$"#,##0'
print(out)
print(s.groupby("pattern").agg(listings=("listing", "count"), nights=("nights_below_min", "sum"),
                                gap=("total_gap_to_min", "sum")))
print(s[["listing", "pattern", "wh_min_price", "wh_adjustment", "nights_below_min", "pct_below",
         "median_price_below", "total_gap_to_min", "median_days_out_below"]].to_string(index=False))
