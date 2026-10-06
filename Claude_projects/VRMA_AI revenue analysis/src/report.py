"""Turn data/scan_targets.pkl into the deliverables:
  output/pricing_flags_<date>.xlsx  every unbooked night (band + flag) + flags + method
  output/pricing_flags_<date>.pdf   top-10 upside / downside tables
"""
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"
RUN = date(2026, 10, 6)
COLOR = {"BLUNDER": "E53935", "MISTAKE": "FB8C00", "INACCURACY": "FDD835"}
STEP_NAME = {1: "1 · same date LY", 2: "2 · comparable dates", 3: "3 · portfolio siblings",
             3.5: "3b · comp set LY", 4: "4 · market only", 5: "5 · skipped"}


def why(r) -> str:
    """One-sentence, committed call."""
    p = f"${r.price:,.0f}"
    if r.side == "DOWN":
        verdict = "Underpriced — raise it into the band."
        if r.step == 1:
            base = f"Asking {p} vs ~${r.floor:,.0f}/night this unit booked on {r.ly_date:%a %b %-d} LY"
        elif r.step == 2:
            base = f"Asking {p} vs ${r.floor:,.0f}, the median this unit booked on comparable LY {r.ly_date:%a}s"
        elif r.step == 3:
            base = f"Asking {p} vs ${r.floor:,.0f}, what size-scaled siblings booked on {r.ly_date:%b %-d} LY (weak comparison)"
        elif r.step == 3.5:
            base = f"Asking {p} vs ${r.floor:,.0f}, what its Wheelhouse comp set booked around {r.ly_date:%b %-d} LY, scaled to this unit (weak comparison)"
        else:
            base = f"Asking {p}, under the bottom of the Wheelhouse comp band (${r.floor:,.0f}; market-only, weak)"
    else:
        verdict = "Overpriced — bring it down into the band."
        if r.step == 1:
            base = f"Asking {p} vs ~${r.floor:,.0f}/night it booked {r.ly_date:%a %b %-d} LY; ceiling ${r.ceiling:,.0f}"
        elif r.step == 2:
            base = f"Asking {p} vs ${r.floor:,.0f} median on comparable LY {r.ly_date:%a}s; ceiling ${r.ceiling:,.0f}"
        elif r.step == 3:
            base = f"Asking {p} vs ${r.floor:,.0f} sibling-implied LY rate (weak comparison)"
        elif r.step == 3.5:
            base = f"Asking {p} vs ${r.ceiling:,.0f} ceiling from its comp set's LY rate, scaled to this unit (weak comparison)"
        else:
            base = f"Asking {p}, above the top of the Wheelhouse comp band (${r.ceiling:,.0f}; market-only, weak)"
    extra = []
    n = r.notes or ""
    if "lifted to BLUNDER" in n:
        extra.append("lifted to BLUNDER: " + n.split("lifted to BLUNDER: ")[1].split(";")[0])
    if "capped at MISTAKE" in n and "holiday" in n:
        extra.append("LY stay spanned a holiday, capped at MISTAKE")
    if r.ly_stay and isinstance(r.ly_stay, str) and r.step == 1:
        nn = int(r.ly_stay.split("(")[1].split("n,")[0])
        if nn >= 7:
            extra.append(f"LY was a {nn}-night discounted stay")
    s = base + (" (" + "; ".join(extra) + ")" if extra else "") + ". " + verdict
    return s


def tidy(df):
    df = df.copy()
    df["why"] = [why(r) if r.severity else "" for r in df.itertuples()]
    df["evidence_step"] = df.step.map(STEP_NAME)
    df["pct_outside"] = df.pct_outside.round(4)
    for c in ("floor", "ceiling"):
        df[c] = df[c].round(0)
    df["was"] = "n/a"
    return df


def write_xlsx(df, path):
    cols = ["listing", "date", "dow", "days_out", "price", "was", "floor", "ceiling", "side",
            "pct_outside", "severity", "evidence_step", "why", "evidence", "ly_date", "ly_match",
            "ly_stay", "market_ratio", "holiday", "event_ly", "near_event", "max_severity",
            "raw_severity", "notes"]
    flags = df[df.severity.notna()].sort_values(["severity", "pct_outside"], key=lambda s: s.map(
        {"BLUNDER": 0, "MISTAKE": 1, "INACCURACY": 2}) if s.name == "severity" else -s)
    method = pd.DataFrame({"item": [
        "Run date", "Horizon", "Scope", "Current price", "Was price", "Floor", "Ceiling",
        "Market + pacing adjustment", "LY night value", "Mid-term stays", "Distressed sale",
        "Holiday cap", "Step 2", "Step 3", "Step 4", "Severity caps", "Corroboration",
        "Holidays", "Events"],
        "rule": [
        RUN.isoformat(), "2026-10-06 .. 2027-04-03 (180 nights), Guesty status = available only",
        "Elektra 1115, Yacinde E3, Shoreline 15510, Longbranch 6821 (whole portfolio banded for sibling evidence)",
        "Guesty calendar nightly price pulled 2026-10-06 (= Wheelhouse recommendation on nearly every night)",
        "Not available — Guesty/Wheelhouse expose no price history; needs a second snapshot to diff",
        "Confirmed, paid (non-owner) LY booking on the matching night; LY = same weekday 364d back, calendar date for fixed holidays, Easter-aligned for Easter",
        "Floor x 1.30",
        "The listing's Wheelhouse comp set: booked ADR +/-7d vs LY where the set has >=15% on the books, otherwise the set's trailing 60-day YoY ADR (clipped 0.75-1.30); whole-market ADR only if no comp set. <90% lowers the floor, >110% lifts the ceiling. Neighborhood pacing (booked vs typical at this lead time, +/-3 days, >=20 expected bookings): <=75% lowers the floor 10%, >=125% lifts the ceiling 10%. Total lift capped at 1.35x",
        "Stay's accommodation fare (fareAccommodationAdjusted, excl. cleaning/taxes) split across nights by the market's daily ADR profile LY",
        "Stays of 28+ nights excluded as evidence (monthly rates)",
        "Booked <=3 days before arrival AND < 60% of siblings' same-night value -> not a floor, drop to step 3",
        "LY stay included a holiday night and the target date is not a holiday -> max MISTAKE",
        "Same listing, same weekday, +/-4 weeks of the LY date, no holidays/events, >=2 nights; median. Not used for holiday/event targets",
        "Siblings = same Wheelhouse comp set with LY history, else same market +/-1 bedroom; each sibling's LY night value scaled by the pair's LY-rate ratio (or Wheelhouse rec-price ratio if the target has no LY); median. Tried before step 2 when the date sits within 2 days of an LY event",
        "Step 3b: the listing's Wheelhouse comp set, median ADR over LY date +/-7d, times how this listing books relative to that set (median of its own booked night / set ADR, >=10 nights). Step 4: Wheelhouse neighborhood comp band, floor = low_price, ceiling = high_price (asking prices, so it runs high)",
        "Steps 3, 3b and 4 max INACCURACY",
        "Step-3 flag >= 30% becomes BLUNDER when >= 2 siblings (same pool) have step-1 BLUNDERs that date on the same side",
        "Halloween, Thanksgiving Wed-Sat, Dec 23-26, Dec 30-Jan 1, MLK & Presidents Sat-Sun, Feb 13-14, Easter Fri-Sat",
        "LY market nights with ADR >= 25% above the same-weekday median of the surrounding 8 weeks"]})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        flags[cols].to_excel(xw, sheet_name="flags", index=False)
        df.sort_values(["listing", "date"])[cols].to_excel(xw, sheet_name="all_unbooked_nights", index=False)
        method.to_excel(xw, sheet_name="method", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
                w = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(9, w + 2))
            hdr = [c.value for c in ws[1]]
            if "severity" in hdr:
                si, pi = hdr.index("severity") + 1, hdr.index("pct_outside") + 1
                for row in ws.iter_rows(min_row=2):
                    sev = row[si - 1].value
                    if sev in COLOR:
                        row[si - 1].fill = PatternFill("solid", fgColor=COLOR[sev])
                        row[si - 1].font = Font(bold=True, color="FFFFFF" if sev != "INACCURACY" else "000000")
                    row[pi - 1].number_format = "0.0%"
                    for k in ("price", "floor", "ceiling"):
                        row[hdr.index(k)].number_format = '"$"#,##0'
            if ws.title == "method":
                ws.column_dimensions["B"].width = 140
                for row in ws.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical="top")


def shared_causes(top):
    lines = []
    for lst, g in top.groupby("listing"):
        if len(g) >= 3:
            steps = g.step.value_counts()
            lines.append(f"{lst}: {len(g)} of these rows — "
                         + ", ".join(f"{n} via step {s}" for s, n in steps.items()))
    return lines


def write_pdf(df, path, notes_up, notes_down):
    st = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=st["BodyText"], fontSize=7, leading=8.5)
    h = ParagraphStyle("h", parent=st["Heading2"], fontSize=12, spaceAfter=4)
    small = ParagraphStyle("s", parent=st["BodyText"], fontSize=7.5, leading=9.5)
    doc = SimpleDocTemplate(str(path), pagesize=landscape(letter), leftMargin=0.35 * inch,
                            rightMargin=0.35 * inch, topMargin=0.4 * inch, bottomMargin=0.4 * inch)
    story = [Paragraph("Pricing errors — next 180 days (unbooked nights)", st["Title"]),
             Paragraph(f"Run {RUN:%b %-d, %Y} · Elektra 1115, Yacinde E3, Shoreline 15510, Longbranch 6821 · "
                       "current price = Guesty calendar · LY bookings = Guesty confirmed · market = Wheelhouse", small),
             Spacer(1, 6)]
    head = ["Listing", "Date", "Days\nout", "Current\n(was)", "Floor", "Ceiling", "% out", "Severity",
            "Evidence step", "Why"]
    widths = [0.95, 0.72, 0.38, 0.58, 0.48, 0.5, 0.42, 0.88, 0.92, 4.42]
    for title, side, notes in (("Top 10 UPSIDE — priced over the ceiling", "UP", notes_up),
                               ("Top 10 DOWNSIDE — priced under the floor", "DOWN", notes_down)):
        top = df[(df.side == side) & df.severity.notna()].sort_values("pct_outside", ascending=False).head(10)
        rows = [head]
        for r in top.itertuples():
            rows.append([r.listing, f"{r.date:%a %b %-d}", str(r.days_out), f"${r.price:,.0f}\n(n/a)",
                         f"${r.floor:,.0f}", f"${r.ceiling:,.0f}", f"{r.pct_outside:.0%}", r.severity,
                         Paragraph(STEP_NAME[r.step], cell), Paragraph(r.why, cell)])
        t = Table(rows, colWidths=[w * inch for w in widths], repeatRows=1)
        style = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5), ("FONT", (0, 1), (-1, -1), "Helvetica", 7.5),
                 ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#263238")),
                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0BEC5")),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7F8")])]
        for i, r in enumerate(top.itertuples(), 1):
            style += [("BACKGROUND", (7, i), (7, i), colors.HexColor("#" + COLOR[r.severity])),
                      ("TEXTCOLOR", (7, i), (7, i), colors.white if r.severity != "INACCURACY" else colors.black),
                      ("FONT", (7, i), (7, i), "Helvetica-Bold", 7.5)]
        t.setStyle(TableStyle(style))
        story += [Paragraph(title, h), t, Spacer(1, 4)]
        story += [Paragraph("• " + n, small) for n in notes]
        story.append(Spacer(1, 10))
    story.append(Paragraph(
        "Was price: not available — neither Guesty nor Wheelhouse exposes price history; a second pull is needed to diff. "
        "Severity colours: BLUNDER ≥30% outside (red), MISTAKE 20–30% (orange), INACCURACY 10–20% (yellow). "
        "Steps 3–4 capped at INACCURACY unless lifted by sibling corroboration. Full list: pricing_flags spreadsheet.", small))
    doc.build(story)


def main(notes_up=(), notes_down=()):
    OUT.mkdir(exist_ok=True)
    df = tidy(pd.read_pickle(ROOT / "data" / "scan_targets.pkl"))
    df.to_pickle(ROOT / "data" / "scan_targets_tidy.pkl")
    tag = RUN.isoformat()
    write_xlsx(df, OUT / f"pricing_flags_{tag}.xlsx")
    write_pdf(df, OUT / f"pricing_flags_{tag}.pdf", notes_up, notes_down)
    print("wrote", OUT / f"pricing_flags_{tag}.xlsx", OUT / f"pricing_flags_{tag}.pdf")


if __name__ == "__main__":
    import json
    import sys
    n = json.loads(Path(sys.argv[1]).read_text()) if len(sys.argv) > 1 else {}
    main(n.get("up", []), n.get("down", []))
