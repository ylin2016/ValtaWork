"""Turn cached extractions into the OSBR expense report. No API calls — rerun freely.

    python -m src.build            # -> output/osbr_expenses.xlsx + output/osbr_items.csv

Tax, shipping/fees and order-level discounts are spread over each receipt's items in
proportion to their amount; pennies are rounded per line and the remainder goes to the
largest item, so each receipt's item totals add up exactly. A receipt whose printed
total doesn't equal items + tax + shipping + fees - discount is flagged.

Cottage: the item's own cottage, else the receipt's (both read by Claude from the document or
path), else "OSBR 5" / "Cottage 5" in the file path, else "Unassigned".

Corrections: add rows to config/overrides.csv
    receipt_file,line_no,category,cottage,note
(receipt_file and line_no as shown on the Items sheet; leave category or cottage blank to keep
it; category "Exclude" drops a personal/non-OSBR line together with its share of tax), then
rerun build.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .paths import (CATEGORIES_YML, EXCLUDED_CSV, EXTRACTED, MANUAL_CSV, OUTPUT, OVERRIDES_CSV,
                    REPORT_MONTHS, STATEMENT_CSV)

CENT = Decimal("0.01")
TOLERANCE = Decimal("0.02")
COTTAGES = [str(i) for i in range(1, 13)]
NAME_AMOUNT = re.compile(r"_\s*\$?(-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?)(?:\.[A-Za-z]{2,4})?$")
# 'Cottage 1_5' = cottages 1 and 5; a number followed by '.digit' is a date ('OSBR 5_7.28')
PATH_COTTAGE = re.compile(r"(?:osbr|cottage)[\s_#-]*((?:\d{1,2}(?![\d.])[\s_&,+]*)+)", re.I)


def d(x) -> Decimal:
    return Decimal(str(x or 0)).quantize(CENT, ROUND_HALF_UP)


def allocate(amounts: list[Decimal], extra: Decimal) -> list[Decimal]:
    """Split `extra` across lines pro rata to amount; rounded shares sum exactly to extra."""
    if not amounts or extra == 0:
        return [Decimal("0.00")] * len(amounts)
    base = sum(amounts)
    if base == 0:
        shares = [Decimal("0.00")] * len(amounts)
    else:
        shares = [(extra * a / base).quantize(CENT, ROUND_HALF_UP) for a in amounts]
    largest = max(range(len(amounts)), key=lambda i: abs(amounts[i]))
    shares[largest] += extra - sum(shares)
    return shares


def norm_cottage(v: str) -> str:
    """'5', 'Cottage 5', 'OSBR 05' -> '5'; 'common' -> 'Common'; else ''."""
    v = (v or "").strip()
    if v.lower() in {"common", "resort", "all"}:
        return "Common"
    m = re.search(r"\d{1,2}", v)
    return str(int(m.group())) if m and str(int(m.group())) in COTTAGES else ""


def cottage_from_path(path: str) -> str:
    """One cottage named in the path -> '5'; none, or several (a shared purchase) -> ''."""
    found = {str(int(n)) for grp in PATH_COTTAGE.findall(path) for n in re.findall(r"\d+", grp)}
    found &= set(COTTAGES)
    return found.pop() if len(found) == 1 else ""


def name_amount(path: str) -> Decimal | None:
    """Trailing amount in the file name, e.g. '..._Housekeeping supplies_304.78' -> 304.78."""
    m = NAME_AMOUNT.search(path.rsplit("/", 1)[-1])
    if not m:
        return None
    raw = m.group(1).replace(",", "")
    # '..._250_7500.pdf' in folder 'Valta Homes - 7500': trailing number is the card, not an amount
    if raw in re.findall(r"\d{4}", path.rsplit("/", 1)[0]):
        return None
    return Decimal(raw)


def name_amounts(path: str) -> set[Decimal]:
    """Every amount in the trailing '_a_b_c' chain of the file name. Split receipts are named
    '<total>_<part1>_<part2>' (e.g. 'picnic area 70.48_17.47_53.01'), so the total may not be last."""
    stem = re.sub(r"\.[A-Za-z]{2,4}$", "", path.rsplit("/", 1)[-1])
    m = re.search(r"((?:[_ ]\$?-?\d+(?:\.\d{1,2})?)+)$", stem)
    return {Decimal(x) for x in re.findall(r"-?\d+(?:\.\d{1,2})?", m.group(1))} if m else set()


def card_from_path(path: str) -> str:
    """Card last 4 from the top folder: '9967/…', 'Amazon 4783/…', 'Credit - 3104/…'."""
    if "/" not in path:
        return ""
    m = re.search(r"(?<!\d)(\d{4})(?!\d)", path.split("/", 1)[0])
    return m.group(1) if m else ""


def default_statement_month(src: str, receipt_date: str) -> str:
    """Owner-statement month a receipt is charged on: the upload date that starts the file
    name ('20260804_OSBR_…' -> 2026-08), else the receipt's own month."""
    # allow a leading tag such as '(refunded) 20260829_…'
    m = re.match(r"(?:\([^)]*\)\s*)?(20\d\d)(0[1-9]|1[0-2])\d\d", src.rsplit("/", 1)[-1])
    return f"{m.group(1)}-{m.group(2)}" if m else receipt_date[:7]


def load_statement_map() -> dict[tuple[str, int | None], str]:
    """config/statement_months.csv: source_file,line_no,statement_month,note.
    Blank line_no = whole file; a line_no moves just that line (a receipt split across statements)."""
    if not STATEMENT_CSV.exists():
        return {}
    out = {}
    with open(STATEMENT_CSV, newline="") as f:
        for r in csv.DictReader(f):
            if (r.get("source_file") or "").strip() and (r.get("statement_month") or "").strip():
                ln = (r.get("line_no") or "").strip()
                out[(r["source_file"].strip(), int(ln) if ln else None)] = r["statement_month"].strip()
    return out


def load_overrides() -> dict[tuple[str, int], dict]:
    if not OVERRIDES_CSV.exists():
        return {}
    out = {}
    with open(OVERRIDES_CSV, newline="") as f:
        for r in csv.DictReader(f):
            if r.get("receipt_file") and (r.get("line_no") or "").strip():
                out[(r["receipt_file"].strip(), int(r["line_no"]))] = {
                    k: (r.get(k) or "").strip() for k in ("category", "cottage", "note")}
    return out


def load_manual() -> list[dict]:
    """config/manual_expenses.csv rows as one-item receipts, source_file 'manual/<source>/<n>'."""
    if not MANUAL_CSV.exists():
        return []
    out = []
    with open(MANUAL_CSV, newline="") as f:
        for n, r in enumerate(csv.DictReader(f), 1):
            if not (r.get("amount") or "").strip():
                continue
            amt = float(r["amount"])
            src = f"manual/{(r.get('source') or 'manual').strip()}/{n}"
            item = {"description": r["description"], "qty": 1, "unit_price": amt, "amount": amt,
                    "category": r["category"].strip(), "cottage": (r.get("cottage") or "").strip(),
                    "confidence": "high", "notes": (r.get("note") or "").strip()}
            receipt = {"vendor": r["vendor"], "date": r["date"].strip(), "document_type": "other",
                       "receipt_number": "", "payment_method": "", "cottage": item["cottage"],
                       "subtotal": amt, "tax": 0, "shipping": 0, "other_fees": 0,
                       "order_discount": 0, "total": amt, "total_legible": True,
                       "items": [item], "notes": "manual entry (no receipt file)"}
            out.append({"source_file": src, "hash": f"manual-{n}", "extracted_at": "",
                        "statement_month": (r.get("statement_month") or "").strip() or r["date"].strip()[:7],
                        "result": {"receipts": [receipt], "not_a_receipt_reason": ""}})
    return out


def load_excluded() -> dict[str, str]:
    if not EXCLUDED_CSV.exists():
        return {}
    with open(EXCLUDED_CSV, newline="") as f:
        return {r["source_file"].strip(): r.get("reason", "") for r in csv.DictReader(f)
                if (r.get("source_file") or "").strip()}


def load_records() -> list[dict]:
    excluded = load_excluded()
    recs = [json.loads(p.read_text()) for p in sorted(EXTRACTED.glob("*.json"))]
    kept = [r for r in recs if r["source_file"] not in excluded]
    missing = set(excluded) - {r["source_file"] for r in recs}
    print(f"{len(recs) - len(kept)} files excluded via {EXCLUDED_CSV.name}"
          + (f"; {len(missing)} listed there not found: {sorted(missing)}" if missing else ""))
    return kept + load_manual()


def build_rows(records, categories, overrides, stmt_map):
    items, receipts, flags, files = [], [], [], []
    seen_totals: dict[tuple, str] = {}
    used_overrides = set()
    used_stmt = set()

    for rec in sorted(records, key=lambda r: r["source_file"]):
        src = rec["source_file"]
        res = rec.get("result", {})
        rs = res.get("receipts", [])
        f_stmt = (stmt_map.get((src, None)) or rec.get("statement_month")
                  or default_statement_month(src, rs[0]["date"] if rs else ""))
        files.append({"source_file": src, "statement_month": f_stmt, "receipts": len(rs),
                      "note": res.get("not_a_receipt_reason", ""),
                      "extracted_at": rec.get("extracted_at", ""), "hash": rec["hash"]})
        if not rs:
            flags.append({"source_file": src, "receipt": "", "issue": "not a receipt",
                          "detail": res.get("not_a_receipt_reason", "")})
        named = name_amount(src)
        read_total = sum(d(r["total"]) for r in rs)
        chain = name_amounts(src)
        if (named is not None and rs and abs(named - read_total) > TOLERANCE
                and not any(abs(a - read_total) <= TOLERANCE for a in chain)):
            flags.append({"source_file": src, "receipt": rs[0]["vendor"], "issue": "file name amount differs",
                          "detail": f"name says {named}, receipt total read as {read_total}"})

        line_no = 0  # numbered across the whole file so overrides stay stable
        for ri, r in enumerate(rs, 1):
            rid = f"{src}#{ri}" if len(rs) > 1 else src
            r_cottage = norm_cottage(r.get("cottage", "")) or cottage_from_path(src)
            if (src, None) in stmt_map:
                used_stmt.add((src, None))
            r_stmt = (stmt_map.get((src, None)) or rec.get("statement_month")
                      or default_statement_month(src, r["date"]))
            first_item = len(items)
            amounts = [d(it["amount"]) for it in r["items"]]
            tax, ship, fees, disc, total = (d(r["tax"]), d(r["shipping"]), d(r["other_fees"]),
                                            d(r["order_discount"]), d(r["total"]))
            a_tax = allocate(amounts, tax)
            a_ship = allocate(amounts, ship + fees)
            a_disc = allocate(amounts, -disc)
            computed = sum(amounts) + tax + ship + fees - disc
            diff = total - computed
            ok = abs(diff) <= TOLERANCE and r["total_legible"]

            if not r["total_legible"]:
                flags.append({"source_file": rid, "receipt": r["vendor"], "issue": "total not legible",
                              "detail": f"read as {total}"})
            elif not ok:
                flags.append({"source_file": rid, "receipt": r["vendor"], "issue": "does not reconcile",
                              "detail": f"printed total {total} vs items+tax+fees {computed} (diff {diff})"})
            if not r["items"]:
                flags.append({"source_file": rid, "receipt": r["vendor"], "issue": "no items read", "detail": ""})
            if not r["date"]:
                flags.append({"source_file": rid, "receipt": r["vendor"], "issue": "no date", "detail": ""})

            key = (r["vendor"].strip().lower(), r["date"], total)
            if key in seen_totals:
                flags.append({"source_file": rid, "receipt": r["vendor"], "issue": "possible duplicate",
                              "detail": f"same vendor/date/total as {seen_totals[key]}"})
            else:
                seen_totals[key] = rid

            receipts.append({
                "source_file": rid, "date": r["date"], "vendor": r["vendor"],
                "document_type": r["document_type"], "receipt_number": r["receipt_number"],
                "payment_method": r["payment_method"], "card_folder": card_from_path(src),
                "name_amount": named if named is not None else "",
                "cottage": r_cottage or "Unassigned",
                "items": len(r["items"]),
                "items_subtotal": sum(amounts), "tax": tax, "shipping_fees": ship + fees,
                "discount": -disc, "total": total, "reconciles": "yes" if ok else "NO",
                "diff": diff, "excluded": Decimal("0.00"), "notes": r["notes"],
            })

            for it, amt, t, s, dc in zip(r["items"], amounts, a_tax, a_ship, a_disc):
                line_no += 1
                cat, conf, note = it["category"], it["confidence"], it["notes"]
                cottage = norm_cottage(it.get("cottage", "")) or r_cottage or "Unassigned"
                ov = overrides.get((src, line_no))
                if ov:
                    used_overrides.add((src, line_no))
                    if ov["category"]:
                        cat, conf = ov["category"], "high"
                    if ov["cottage"]:
                        cottage = norm_cottage(ov["cottage"]) or ov["cottage"]
                    note = f"[override] {ov['note']}".strip()
                if cat.lower() == "exclude":  # personal / non-OSBR line: drop it with its tax share
                    receipts[-1]["excluded"] += amt + t + s + dc
                    continue
                if cat not in categories:
                    flags.append({"source_file": src, "receipt": r["vendor"], "issue": "unknown category",
                                  "detail": f"line {line_no}: {cat}"})
                if conf == "low":
                    flags.append({"source_file": src, "receipt": r["vendor"], "issue": "low confidence",
                                  "detail": f"line {line_no}: {it['description']} -> {cat}. {note}"})
                stmt = stmt_map.get((src, line_no), r_stmt)
                if (src, line_no) in stmt_map:
                    used_stmt.add((src, line_no))
                items.append({
                    "source_file": src, "line_no": line_no, "date": r["date"],
                    "month": r["date"][:7], "statement_month": stmt,
                    "vendor": r["vendor"], "cottage": cottage,
                    "description": it["description"],
                    "qty": it["qty"], "unit_price": d(it["unit_price"]), "amount": amt,
                    "allocated_tax": t, "allocated_shipping_fees": s, "allocated_discount": dc,
                    "total": amt + t + s + dc, "category": cat, "confidence": conf, "notes": note,
                })

            receipts[-1]["statement_month"] = " + ".join(
                sorted({it["statement_month"] for it in items[first_item:]})) or r_stmt

    for k in stmt_map.keys() - used_stmt:
        flags.append({"source_file": k[0], "receipt": "", "issue": "statement month not applied",
                      "detail": f"line {k[1]} not found" if k[1] else "file not found"})
    for k in overrides.keys() - used_overrides:
        flags.append({"source_file": k[0], "receipt": "", "issue": "override not applied",
                      "detail": f"line {k[1]} not found"})
    return items, receipts, flags, files


# ---------- output ----------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0.00;[Red]-#,##0.00'


def write_sheet(ws, rows: list[dict], columns: list[str], money: set[str] = frozenset()):
    ws.append(columns)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        ws.append([float(r[c]) if c in money else r.get(c, "") for c in columns])
    for i, col in enumerate(columns, 1):
        letter = get_column_letter(i)
        if col in money:
            for cell in ws[letter][1:]:
                cell.number_format = MONEY
        width = max([len(col)] + [len(str(r.get(col, ""))) for r in rows[:500]])
        ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def write_summary(ws, items, receipts, flags, categories, months):
    by_cat = defaultdict(lambda: [0, Decimal(0)])
    by_stmt = defaultdict(lambda: defaultdict(Decimal))
    by_cottage = defaultdict(lambda: defaultdict(Decimal))
    for it in items:
        by_cat[it["category"]][0] += 1
        by_cat[it["category"]][1] += it["total"]
        by_stmt[it["statement_month"] or "none"][it["category"]] += it["total"]
        by_cottage[it["cottage"]][it["category"]] += it["total"]
    grand = sum(v[1] for v in by_cat.values())

    ws.append([f"OSBR expenses by category — owner statements {', '.join(months)}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{len(receipts)} receipts, {len(items)} items, "
               f"{sum(1 for r in receipts if r['reconciles'] == 'NO')} not reconciling, "
               f"{len(flags)} flags. Built {date.today()}."])
    ws.append([])
    ws.append(["Category", "Items", "Total", "% of spend"])
    for c in ws[4]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    order = [c for c in categories if c in by_cat] + [c for c in by_cat if c not in categories]
    for cat in sorted(order, key=lambda c: -by_cat[c][1]):
        n, tot = by_cat[cat]
        ws.append([cat, n, float(tot), float(tot / grand) if grand else 0])
    ws.append(["Total", len(items), float(grand), 1 if grand else 0])
    last = ws.max_row
    for row in ws.iter_rows(min_row=5, max_row=last):
        row[2].number_format, row[3].number_format = MONEY, "0.0%"
    for c in ws[last]:
        c.font = Font(bold=True)

    cats = sorted(order, key=lambda c: -by_cat[c][1])
    cottage_order = COTTAGES + ["Common", "Unassigned"]
    cottage_keys = sorted(by_cottage, key=lambda k: (cottage_order.index(k) if k in cottage_order else 99, k))
    pivot(ws, "Statement month", [(m, by_stmt[m]) for m in sorted(by_stmt)], cats)
    pivot(ws, "Cottage", [(f"Cottage {k}" if k in COTTAGES else k, by_cottage[k]) for k in cottage_keys], cats)
    ws.column_dimensions["A"].width = 18
    for i in range(2, len(cats) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws["A2"].alignment = Alignment(wrap_text=False)


def pivot(ws, label: str, rows: list[tuple[str, dict]], cats: list[str]) -> None:
    ws.append([])
    start = ws.max_row + 1
    ws.append([label] + cats + ["Total"])
    for c in ws[start]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    col_tot = defaultdict(Decimal)
    for name, vals in rows:
        v = [vals.get(c, Decimal(0)) for c in cats]
        for c, x in zip(cats, v):
            col_tot[c] += x
        ws.append([name] + [float(x) for x in v] + [float(sum(v))])
    ws.append(["Total"] + [float(col_tot[c]) for c in cats] + [float(sum(col_tot.values()))])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=start + 1, max_row=ws.max_row, min_col=2):
        for c in row:
            c.number_format = MONEY


def in_scope(rows: list[dict], months: set[str], key: str = "statement_month") -> list[dict]:
    return [r for r in rows if set(str(r.get(key, "")).split(" + ")) & months]


def main() -> None:
    ap = argparse.ArgumentParser(description="Build the OSBR expense report (no API calls).")
    ap.add_argument("--months", help=f"comma-separated statement months (default {','.join(REPORT_MONTHS)})")
    args = ap.parse_args()
    months = sorted(args.months.split(",") if args.months else REPORT_MONTHS)
    categories = list(yaml.safe_load(CATEGORIES_YML.read_text())["categories"])
    records = load_records()
    if not records:
        raise SystemExit("No extractions in output/extracted/. Run: python -m src.extract")
    items, receipts, flags, files = build_rows(records, categories, load_overrides(),
                                               load_statement_map())
    scope = set(months)
    file_month = {f["source_file"]: f["statement_month"] for f in files}
    items, receipts, files = in_scope(items, scope), in_scope(receipts, scope), in_scope(files, scope)
    # flag source_file is the file, or "<file>#<n>" for one receipt of several
    flags = [f for f in flags if (file_month.get(f["source_file"])
                                  or file_month.get(f["source_file"].rsplit("#", 1)[0])) in scope]

    item_cols = ["source_file", "line_no", "statement_month", "date", "month", "vendor", "cottage", "description", "qty",
                 "unit_price", "amount", "allocated_tax", "allocated_shipping_fees",
                 "allocated_discount", "total", "category", "confidence", "notes"]
    item_money = {"unit_price", "amount", "allocated_tax", "allocated_shipping_fees",
                  "allocated_discount", "total"}
    rec_cols = ["source_file", "statement_month", "date", "vendor", "cottage", "document_type", "receipt_number",
                "payment_method", "card_folder", "name_amount", "items", "items_subtotal", "tax", "shipping_fees",
                "discount", "total", "reconciles", "diff", "excluded", "notes"]
    rec_money = {"items_subtotal", "tax", "shipping_fees", "discount", "total", "diff", "excluded"}

    items.sort(key=lambda r: (r["statement_month"], r["date"], r["source_file"], r["line_no"]))
    receipts.sort(key=lambda r: (r["statement_month"], r["date"], r["source_file"]))

    wb = Workbook()
    write_summary(wb.active, items, receipts, flags, categories, months)
    wb.active.title = "Summary"
    write_sheet(wb.create_sheet("Items"), items, item_cols, item_money)
    write_sheet(wb.create_sheet("Receipts"), receipts, rec_cols, rec_money)
    write_sheet(wb.create_sheet("Flags"), flags, ["source_file", "receipt", "issue", "detail"])
    write_sheet(wb.create_sheet("Files"), files, ["source_file", "statement_month", "receipts", "note",
                                                  "extracted_at", "hash"])

    OUTPUT.mkdir(exist_ok=True)
    xlsx = OUTPUT / "osbr_expenses.xlsx"
    wb.save(xlsx)
    with open(OUTPUT / "osbr_items.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=item_cols)
        w.writeheader()
        w.writerows(items)

    total = sum(it["total"] for it in items)
    print(f"{len(files)} files -> {len(receipts)} receipts -> {len(items)} items, total ${total:,.2f}")
    print(f"{sum(1 for r in receipts if r['reconciles'] == 'NO')} receipts don't reconcile; "
          f"{len(flags)} flags (see Flags sheet)")
    print(f"Wrote {xlsx}")


if __name__ == "__main__":
    main()
