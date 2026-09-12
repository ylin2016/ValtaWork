"""Build QuickBooks Expenses from the Zelle payment tracking workbook.

One Zelle transfer pays several invoices at once.  The workbook records the
invoices, one row each; the `Zelle content` column is what ties the rows of a
single transfer together, and `JE Amount` is that transfer's total.

    Expense   paid from Chase Trust Checking 9967 - STR, vendor = the payee
              one line per invoice -> owner-expense account, classed to the property

This is an Expense (`Purchase`), not a journal entry: it is the same shape the
hand-entered ones already have (see Purchase 116690, Andrea Brannon 2026-08-29,
$1,487.69 over 13 lines), so the bank feed matches it the way it matches those.

    python -m src.invoices.build_zelle                       # -> review/zelle_expenses.csv
    python -m src.invoices.build_zelle --skip-existing       # drop transfers already booked

Category -> account comes from the workbook's own `account_map` sheet when it has one,
else from `config/Account Mapping.csv` (or `--map`), rather than from
`config/accounts.yml`: the owner maintains it alongside the data, and every name is
resolved against the live company file, so a typo fails loudly instead of posting to
the wrong account.

Two workbook layouts are read:

    Zelle_payment_tracking.xlsx   sheet `Records`; Amount and Property are columns
    <date>_Expense2qbo.xlsx       first sheet; Amount and Property live only in the
                                  `filename`:  <date>_<Property>_<payee>_..._<amount>_<lump>

    python -m src.invoices.build_zelle --src inputs/Invoice_payment/20260911_Expense2qbo.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

import yaml

from ..config import client, location, name as cfg_name, payees
from ..paths import CONFIG_DIR, INVOICE_INPUTS, review_csv
from ..resolver import Resolver

BANK_ACCT = cfg_name("bank_str")
LOCATION = location()

DEFAULT_IN = INVOICE_INPUTS / "Zelle_payment_tracking.xlsx"

RECORDS_SHEET = "Records"
ACCOUNT_MAP_SHEET = "account_map"
DEFAULT_MAP_CSV = CONFIG_DIR / "Account Mapping.csv"

# Sheet property names that are not a listing.  "All Units" is company-wide supplies,
# and 45 of the 46 lines ever posted to `All units supplies` carry class Valta Realty.
PROPERTY_CLASSES_YML = CONFIG_DIR / "property_classes.yml"
LINE_SPLITS_YML = CONFIG_DIR / "line_splits.yml"

# QuickBooks truncates a DocNumber past this, and a silently shortened reference is
# a reference that no longer matches the bank memo.
DOCNUMBER_MAX = 21


def iso(v) -> str:
    """The Date column, however Excel chose to store it, as YYYY-MM-DD."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):          # an unformatted Excel serial
        return (date(1899, 12, 30) + timedelta(days=float(v))).isoformat()
    return str(v).strip()[:10]


def rows_of(ws) -> list[dict]:
    """Sheet rows as dicts keyed by the header row, blank rows dropped."""
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for raw in it:
        if all(c is None or str(c).strip() == "" for c in raw):
            continue
        out.append({h: v for h, v in zip(header, raw) if h})
    return out


def load_workbook_data(src: Path, map_csv: Path | None = None
                       ) -> tuple[list[dict], dict[str, str], str]:
    """Records, category -> account (keys lower-cased), and where the map came from.

    Category is matched case-insensitively: the sheet says `Cleaning Payout` on some
    rows and `Cleaning payout` on others, and those are one category, not two.
    """
    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb[RECORDS_SHEET] if RECORDS_SHEET in wb.sheetnames else wb.worksheets[0]
    records = rows_of(ws)
    if map_csv is None and ACCOUNT_MAP_SHEET in wb.sheetnames:
        pairs = [(r.get("Categorization"), r.get("qbo_account")) for r in rows_of(wb[ACCOUNT_MAP_SHEET])]
        origin = f"{src.name}:{ACCOUNT_MAP_SHEET}"
    else:
        path = map_csv or DEFAULT_MAP_CSV
        with open(path, encoding="utf-8-sig") as fh:
            pairs = [(r.get("Category") or r.get("Categorization"), r.get("qbo_account"))
                     for r in csv.DictReader(fh)]
        origin = str(path.name)
    wb.close()
    amap = {str(k).strip().lower(): str(v).strip() for k, v in pairs if k and v}
    return records, amap, origin


def property_classes() -> dict[str, str]:
    if not PROPERTY_CLASSES_YML.exists():
        return {}
    return (yaml.safe_load(PROPERTY_CLASSES_YML.read_text()) or {}).get("properties", {}) or {}


def line_splits() -> list[dict]:
    if not LINE_SPLITS_YML.exists():
        return []
    return (yaml.safe_load(LINE_SPLITS_YML.read_text()) or {}).get("splits", []) or []


def desc_of(r: dict) -> str:
    return str(r.get("invoice rename") or r.get("filename") or r.get("InvoiceContent") or "").strip()


def filename_parts(r: dict) -> list[str]:
    """`20260911_Bainbridge 11143_Camila_INV349_9.07_Cleaning_200_605` -> its fields.

    The layout the owner names files in: date, property, payee, ..., line amount, and
    the transfer total last.  Only the ends are fixed; the middle varies by vendor."""
    return [t.strip() for t in desc_of(r).split("_")]


def amount_of(r: dict) -> str | None:
    if r.get("Amount") not in (None, ""):
        return str(r["Amount"])
    parts = filename_parts(r)
    return parts[-2] if len(parts) >= 4 else None


def property_of(r: dict) -> str:
    if r.get("Property"):
        return str(r["Property"]).strip()
    parts = filename_parts(r)
    return parts[1] if len(parts) >= 4 else ""


def compact_docnumber(txndate: str, payee: str, lump: float | None) -> str:
    """`20260908_Andea_455.93` -- date, payee, transfer total, within 21 chars.

    The long `Zelle content` cannot be the DocNumber: truncated to 21 chars,
    `20260911_Camila_Cleaning Inv349-351_605` and `20260911_Camila_Cleaning Payout_...`
    both become `20260911_Camila_Clean`, and the DocNumber duplicate check would then
    refuse the second transfer as already posted.  The full reference goes in the memo.
    """
    amt = "" if lump is None else f"{lump:.2f}".rstrip("0").rstrip(".")
    head, tail = txndate.replace("-", ""), f"_{amt}" if amt else ""
    room = DOCNUMBER_MAX - len(head) - len(tail) - 1
    return f"{head}_{payee.replace(' ', '')[:max(room, 1)]}{tail}"


def build(src: Path, qbo, map_csv: Path | None = None) -> tuple[list[dict], list[str]]:
    records, amap, map_origin = load_workbook_data(src, map_csv)
    print(f"category map: {map_origin} ({len(amap)} categories)")
    aliases = property_classes()
    splits = line_splits()
    res = Resolver(qbo)
    warnings: list[str] = []

    transfers: dict[str, list[dict]] = defaultdict(list)
    for i, r in enumerate(records, start=2):     # +2: header row, 1-based
        ref = str(r.get("Zelle content") or "").strip()
        if not ref:
            warnings.append(f"row {i}: no 'Zelle content' — cannot tell which transfer "
                            f"this invoice was paid by; row DROPPED")
            continue
        r["_row"] = i
        transfers[ref].append(r)

    out: list[dict] = []
    for ref, items in transfers.items():
        txndate = iso(items[0]["Date"])
        if len({iso(r["Date"]) for r in items}) > 1:
            warnings.append(f"{ref}: rows disagree on Date — using {txndate}")

        payto = str(items[0].get("Pay") or items[0].get("Payto") or "").strip()
        vendor = payees().get(payto, payto)
        if not vendor:
            vendor = f"*** NO VENDOR ***"
            warnings.append(f"{ref}: no payee on the row")
        elif res.vendor(vendor) is None:
            warnings.append(f"{ref}: {payto!r} -> {vendor!r} is not a QuickBooks vendor "
                            f"— add it to config/payees.yml")
            vendor = f"*** NO VENDOR {payto} ***"

        # The transfer total the owner recorded, not a sum we derive: if the lines do
        # not add up to it, one is missing or mistyped and nothing should post.
        stated = {round(float(r["JE Amount"]), 2) for r in items if r.get("JE Amount") is not None}
        lump = sorted(stated)[0] if stated else None
        if len(stated) > 1:
            warnings.append(f"{ref}: rows disagree on JE Amount {sorted(stated)} — using {lump}")

        # Which bank the transfer left.  A sheet may say so per row; otherwise the STR
        # trust account, which is where every Zelle so far has come from.
        paid_from = str(items[0].get("Paid from") or items[0].get("Bank") or BANK_ACCT).strip()
        if res.account(paid_from) is None:
            warnings.append(f"{ref}: paid-from account {paid_from!r} is not in this company file")
        docnum = compact_docnumber(txndate, payto or "unknown", lump)

        line = 0
        for r in items:
            line += 1
            amt = amount_of(r)
            if amt is None or str(amt).strip() == "":
                warnings.append(f"{ref} line {line} (row {r['_row']}): no Amount — line DROPPED")
                line -= 1
                continue
            try:
                amt = round(float(amt), 2)
            except ValueError:
                warnings.append(f"{ref} line {line} (row {r['_row']}): amount {amt!r} read from "
                                f"the filename is not a number — line DROPPED")
                line -= 1
                continue
            parts = filename_parts(r)
            if (lump is not None and "Amount" not in r and len(parts) >= 4
                    and _num(parts[-1]) is not None and abs(_num(parts[-1]) - lump) > 0.005):
                warnings.append(f"{ref} line {line}: filename ends in {parts[-1]} but JE Amount is "
                                f"{lump} — the filename layout may not be <amount>_<total>")

            cat = str(r.get("Category") or "").strip()
            acct = amap.get(cat.lower())
            if not acct:
                acct = f"*** NEEDS ACCOUNT {cat or '(blank)'} ***"
                warnings.append(f"{ref} line {line}: category {cat or '(blank)'!r} is not in "
                                f"{map_origin}")
            elif res.account(acct) is None:
                warnings.append(f"{ref} line {line}: account_map points at {acct!r}, which is "
                                f"not an account in this company file")
                acct = f"*** NEEDS ACCOUNT {cat} ***"

            prop = property_of(r)
            cls = ""
            if prop:
                for cand in ([aliases[prop]] if prop in aliases else []) + [f"Listings:{prop}", prop]:
                    if res.klass(cand) is not None:
                        cls = cand
                        break
            if not cls:
                cls = f"*** UNMAPPED {prop or '(blank)'} ***"
                warnings.append(f"{ref} line {line}: no class for property "
                                f"{prop or '(blank)'!r}")

            desc = desc_of(r)
            base = {"PayRef": ref, "DocNumber": docnum, "TxnDate": txndate, "Vendor": vendor,
                    "PaidFrom": paid_from, "Location": LOCATION, "Class": cls,
                    "Category": cat, "Property": prop,
                    "Invoice": str(r.get("Invoice") or r.get("fileName") or "").strip(),
                    "LumpTotal": f"{lump:.2f}" if lump is not None else ""}

            # A line that pays for two things is carved in two; see config/line_splits.yml.
            pieces = [(amt, acct, desc)]
            for sp in splits:
                if not re.search(sp["match"], desc):
                    continue
                cut = round(float(sp["amount"]), 2)
                if amt + 0.005 < cut:
                    warnings.append(f"{ref} line {line}: {sp['name']} line is {amt:.2f}, less than "
                                    f"the {cut:.2f} {sp['name']} charge — NOT split, check it")
                    break
                if res.account(sp["account"]) is None:
                    warnings.append(f"{ref} line {line}: split account {sp['account']!r} is not "
                                    f"in this company file")
                rest = round(amt - cut, 2)
                pieces = [(cut, sp["account"], f"{desc} | {sp['name']} {cut:.2f}")]
                if rest > 0.005:
                    pieces.append((rest, acct, f"{desc} | remainder {rest:.2f}"))
                break

            for i, (a, ac, d) in enumerate(pieces):
                if i:
                    line += 1
                out.append({**base, "LineNum": line, "Amount": f"{a:.2f}",
                            "Account": ac, "Description": d})

        booked = round(sum(float(o["Amount"]) for o in out if o["PayRef"] == ref), 2)
        if lump is None:
            warnings.append(f"{ref}: no JE Amount recorded — lines total {booked:,.2f}")
        elif abs(booked - lump) > 0.005:
            warnings.append(f"{ref}: lines total {booked:,.2f} but JE Amount says "
                            f"{lump:,.2f} — off by {booked - lump:,.2f}")
    docs = defaultdict(list)
    for o in out:
        if o["LineNum"] == 1:
            docs[o["DocNumber"]].append(o["PayRef"])
    for d, refs in docs.items():
        if len(refs) > 1:
            warnings.append(f"DocNumber {d} would be shared by {refs} — same date, payee and "
                            f"total; the second would be refused as a duplicate")
    return out, warnings


def _num(s: str) -> float | None:
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(DEFAULT_IN))
    ap.add_argument("--out", default=str(review_csv("zelle_expenses")))
    ap.add_argument("--map", default=None,
                    help="Category,qbo_account CSV; default: the workbook's account_map sheet, "
                         "else config/Account Mapping.csv")
    ap.add_argument("--skip-existing", action="store_true",
                    help="drop transfers already paid out of the bank account")
    args = ap.parse_args()

    qbo = client()
    rows, warnings = build(Path(args.src), qbo, Path(args.map) if args.map else None)

    if args.skip_existing:
        # Content key (date, amount on the bank account), not the reference: the bank
        # feed books its own Expense for the same Zelle transfer under no reference at
        # all, and a reference check would miss it and pay the invoices twice.
        from .post import existing_zelle
        booked = existing_zelle(qbo)
        totals = defaultdict(float)
        first = {}
        for r in rows:
            totals[r["PayRef"]] += float(r["Amount"])
            first.setdefault(r["PayRef"], r)
        drop = {ref for ref in totals
                if any((first[ref]["PaidFrom"], d, round(totals[ref], 2)) in booked
                       for d in near_dates(first[ref]["TxnDate"]))}
        if drop:
            print(f"skip-existing: {len(drop)} of {len(totals)} already in QBO, "
                  f"{len(totals) - len(drop)} new")
        rows = [r for r in rows if r["PayRef"] not in drop]

    fields = ["PayRef", "DocNumber", "TxnDate", "Vendor", "PaidFrom", "Location", "LineNum", "Amount",
              "Account", "Class", "Description", "Category", "Property", "Invoice", "LumpTotal"]
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

    refs = {r["PayRef"] for r in rows}
    total = sum(float(r["Amount"]) for r in rows)
    print(f"{len(refs)} expense(s), {len(rows)} lines -> {out}")
    print(f"  total {total:,.2f}")
    for w_ in warnings:
        print(f"  WARN {w_}")


def near_dates(d: str, days: int = 3) -> list[str]:
    """The date and its neighbours -- a Zelle transfer can post a day or two late."""
    d0 = datetime.strptime(d, "%Y-%m-%d").date()
    return [(d0 + timedelta(days=n)).isoformat() for n in range(-days, days + 1)]


if __name__ == "__main__":
    main()
