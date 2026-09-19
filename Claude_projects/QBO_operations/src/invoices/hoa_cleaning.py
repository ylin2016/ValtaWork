"""Invoice the Yacinde HOA for the cleans it bears, clearing the receivable.

`fixes/yacinde_hoa_split` parks the HOA's share of each AA Professional Cleaners bill in
`HOA Receivable - Yacinde`.  That balance is money the trust laid out on the HOA's behalf,
and it sits there until the HOA is billed.  This is the other half:

    Bill     Dr HOA Receivable - Yacinde   Cr A/P                 (fixes/yacinde_hoa_split)
    Invoice  Dr Accounts Receivable        Cr HOA Receivable      (here)
    Payment  Dr Chase Trust 9967           Cr Accounts Receivable (the owner, in the UI)

So the invoice item credits the RECEIVABLE, not an income account: recovering a cost is not
revenue, and nothing here should touch the P&L.  Items in this company file may point at a
balance-sheet account -- `Guest Charges:Tax` credits Rental Taxes Payable -- so this is the
house pattern, not a workaround.  A margin, if one is ever charged, is a separate line
against a real income account.

The receivable then nets to zero once everything laid out has been invoiced, which makes it
the control account: a balance means cleans you have paid for and not yet billed on.

Lines carry class `Yacinde HOA`, like the bill lines they clear.  `Owner_statement_whole`'s
`qbo_sync` ingests Invoice lines on the same terms as Bill lines -- any mapped class, no
account filter -- so a listing class here would put HOA cost on an owner's statement.

    python -m src.invoices.hoa_cleaning                          # dry run + review CSV
    python -m src.invoices.hoa_cleaning --confirm                # write
    python -m src.invoices.hoa_cleaning --period 202607          # one month only

The duplicate check is CONTENT-based, per this project's rule: (Customer, TxnDate, TotalAmt)
over existing Invoices, plus the DocNumber.  Neither alone is enough -- a DocNumber is ours
to choose and says nothing about what is already on the books.
"""
from __future__ import annotations

import argparse
import calendar
import csv
import datetime as dt
import sys
from pathlib import Path

import openpyxl

from ..config import client, location, name as acct_name_of
from ..paths import review_csv
from ..qbo_client import QBOClient
from ..resolver import Resolver

CUSTOMER = "Yacinde HOA"
HOA_CLASS = "Yacinde HOA"
ITEM = "HOA Reimbursement - Cleaning"

HOA_SHEET = "Invoice - HOA"
WORKBOOK_DIR = Path(__file__).resolve().parents[3] / "yacinde_expense" / "output"


def workbooks(period: str | None) -> list[Path]:
    pat = f"yacinde_expense_allocation_{period or '2026??'}.xlsx"
    return sorted(WORKBOOK_DIR.glob(pat))


def read_period(path: Path) -> tuple[str, list[dict]]:
    """(YYYY-MM, cleans) off one workbook's `Invoice - HOA` tab.

    The tab's final row is a TOTAL: it is a footer, not a clean, and invoicing it would
    double the bill.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if HOA_SHEET not in wb.sheetnames:
        sys.exit(f"{path.name}: no {HOA_SHEET!r} sheet")
    rows = wb[HOA_SHEET].iter_rows(values_only=True)
    hdr = list(next(rows))
    out, months = [], set()
    for r in rows:
        d = dict(zip(hdr, r))
        unit = str(d.get("unit") or "").strip()
        if not unit or unit.upper() == "TOTAL" or d.get("date") is None:
            continue
        amt = round(float(d.get("amount") or 0), 2)
        if amt <= 0:
            continue
        date = str(d.get("date"))[:10]
        months.add(str(d.get("month") or date[:7]))
        out.append({"Date": date, "Unit": unit, "Amount": amt,
                    "Invoice": str(d.get("aa_invoice") or "").strip(),
                    "Notes": str(d.get("notes") or "").strip()})
    if len(months) != 1:
        sys.exit(f"{path.name}: expected one month, found {sorted(months)}")
    out.sort(key=lambda x: (x["Date"], x["Unit"]))
    return months.pop(), out


def month_end(period: str) -> str:
    y, m = int(period[:4]), int(period[5:7])
    return f"{period}-{calendar.monthrange(y, m)[1]:02d}"


def describe(c: dict) -> str:
    d = c["Date"]
    txt = f"{d[5:7]}/{d[8:10]}/{d[:4]}  {c['Unit']}  cleaning"
    if c["Invoice"]:
        txt += f"  (AA invoice {c['Invoice']})"
    if c["Notes"]:
        txt += f" — {c['Notes']}"
    return txt


def build(period: str, cleans: list[dict], ids: dict, due_days: int) -> dict:
    txndate = month_end(period)
    due = (dt.date.fromisoformat(txndate) + dt.timedelta(days=due_days)).isoformat()
    lines = []
    for n, c in enumerate(cleans, start=1):
        lines.append({
            "LineNum": n,
            "Description": describe(c),
            "Amount": c["Amount"],
            "DetailType": "SalesItemLineDetail",
            "SalesItemLineDetail": {
                "ItemRef": {"value": ids["item"]},
                "ClassRef": {"value": ids["class"]},
                "Qty": 1,
                "UnitPrice": c["Amount"],
                "TaxCodeRef": {"value": "NON"},
            },
        })
    return {
        "CustomerRef": {"value": ids["customer"]},
        "DepartmentRef": {"value": ids["location"]},
        "DocNumber": f"HOA-{period}",
        "TxnDate": txndate,
        "DueDate": due,
        "PrivateNote": (f"Yacinde HOA share of AA Professional Cleaners cleaning, {period}. "
                        f"Clears {acct_name_of('hoa_receivable_yacinde')}."),
        "Line": lines,
    }


def existing(qbo: QBOClient, customer_id: str) -> list[dict]:
    if not customer_id:
        return []
    return qbo.query_all(f"SELECT * FROM Invoice WHERE CustomerRef = '{customer_id}'", "Invoice")


def create_customer(qbo: QBOClient) -> str:
    return qbo.post("customer", {"DisplayName": CUSTOMER, "CompanyName": CUSTOMER})["Id"]


def create_item(qbo: QBOClient, income_acct_id: str) -> str:
    return qbo.post("item", {
        "Name": ITEM, "Type": "Service", "Taxable": False,
        "IncomeAccountRef": {"value": income_acct_id},
        "Description": "Cleaning paid by Valta on the HOA's behalf, rebilled at cost.",
    })["Id"]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--period", default=None, help="YYYYMM; default = every workbook found")
    ap.add_argument("--due-days", type=int, default=30)
    ap.add_argument("--out", default=None)
    ap.add_argument("--create-missing", action="store_true",
                    help="with --confirm, also create the Customer and Item if absent")
    ap.add_argument("--confirm", action="store_true", help="actually WRITE to QuickBooks")
    args = ap.parse_args()

    books = workbooks(args.period)
    if not books:
        sys.exit(f"no workbooks matching period {args.period!r} in {WORKBOOK_DIR}")

    qbo = client()
    res = Resolver(qbo)
    acct = acct_name_of("hoa_receivable_yacinde")
    ids = {
        "customer": res.customer(CUSTOMER),
        "item": res._one("Item", "FullyQualifiedName", ITEM),
        "class": res.klass(HOA_CLASS),
        "location": res.department(location()),
        "account": res.account(acct),
    }
    print("resolving:")
    for k, label in (("customer", f"Customer {CUSTOMER!r}"), ("item", f"Item {ITEM!r}"),
                     ("class", f"Class {HOA_CLASS!r}"), ("location", f"Location {location()!r}"),
                     ("account", f"Account {acct!r}")):
        print(f"  {label:<62} {'Id ' + ids[k] if ids[k] else '*** MISSING ***'}")
    if not ids["account"]:
        sys.exit(f"ABORT — {acct!r} does not exist; run fixes/yacinde_hoa_split first")

    prior = existing(qbo, ids["customer"])
    seen_doc = {i.get("DocNumber") for i in prior}
    seen_key = {(i.get("TxnDate"), round(float(i.get("TotalAmt") or 0), 2)) for i in prior}
    if prior:
        print(f"\n{len(prior)} invoice(s) already on {CUSTOMER}")

    rows, payloads = [], []
    for b in books:
        period, cleans = read_period(b)
        inv = build(period, cleans, ids, args.due_days)
        total = round(sum(c["Amount"] for c in cleans), 2)
        dup = ""
        if inv["DocNumber"] in seen_doc:
            dup = f"DUPLICATE DocNumber {inv['DocNumber']}"
        elif (inv["TxnDate"], total) in seen_key:
            dup = f"DUPLICATE content ({inv['TxnDate']}, {total:,.2f})"
        print(f"\n{b.name}")
        print(f"  {inv['DocNumber']}  {inv['TxnDate']}  due {inv['DueDate']}  "
              f"{len(cleans)} cleans  ${total:,.2f}" + (f"   *** {dup} ***" if dup else ""))
        by_unit: dict[str, list[float]] = {}
        for c in cleans:
            by_unit.setdefault(c["Unit"], []).append(c["Amount"])
        for u in sorted(by_unit):
            print(f"      {u:<14} {len(by_unit[u]):>2} cleans  ${sum(by_unit[u]):>8,.2f}")
        for n, c in enumerate(cleans, start=1):
            rows.append({"DocNumber": inv["DocNumber"], "TxnDate": inv["TxnDate"],
                         "Customer": CUSTOMER, "LineNum": n, "CleanDate": c["Date"],
                         "Unit": c["Unit"], "Amount": f"{c['Amount']:.2f}",
                         "AAInvoice": c["Invoice"], "Item": ITEM, "Class": HOA_CLASS,
                         "CreditsAccount": acct, "Description": describe(c),
                         "InvoiceTotal": f"{total:.2f}", "Duplicate": dup,
                         "Source": b.name})
        payloads.append((inv, total, dup))

    out = Path(args.out) if args.out else review_csv("hoa_cleaning_invoice")
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    grand = sum(t for _, t, _ in payloads)
    print(f"\n{len(payloads)} invoice(s), {len(rows)} lines, ${grand:,.2f} -> {out}")
    print(f"each line debits Accounts Receivable and credits {acct}")

    missing = [k for k in ("customer", "item", "class", "location") if not ids[k]]
    if not args.confirm:
        print("\nDRY RUN — nothing written. Review the CSV, then re-run with --confirm.")
        if missing:
            print(f"\n  create first (or pass --create-missing): {', '.join(missing)}")
        return

    if "class" in missing or "location" in missing:
        sys.exit(f"ABORT — {', '.join(missing)} must exist before posting")
    for k, maker in (("customer", lambda: create_customer(qbo)),
                     ("item", lambda: create_item(qbo, ids["account"]))):
        if ids[k]:
            continue
        if not args.create_missing:
            sys.exit(f"ABORT — {k} {CUSTOMER if k == 'customer' else ITEM!r} does not exist; "
                     f"create it or pass --create-missing")
        ids[k] = maker()
        print(f"created {k}: Id {ids[k]}")

    posted = 0
    for inv, total, dup in payloads:
        if dup:
            print(f"  SKIP {inv['DocNumber']}: {dup}")
            continue
        inv["CustomerRef"]["value"] = ids["customer"]        # may have just been created
        for ln in inv["Line"]:
            ln["SalesItemLineDetail"]["ItemRef"]["value"] = ids["item"]
        got = qbo.post("invoice", inv)
        after = round(float(got.get("TotalAmt") or 0), 2)
        if abs(after - total) > 0.005:
            sys.exit(f"ABORT — {inv['DocNumber']}: posted {after:,.2f}, expected {total:,.2f}")
        posted += 1
        print(f"  Invoice {got['Id']}  {got.get('DocNumber')}  ${after:,.2f}")
    print(f"\nposted {posted} invoice(s)")


if __name__ == "__main__":
    main()
