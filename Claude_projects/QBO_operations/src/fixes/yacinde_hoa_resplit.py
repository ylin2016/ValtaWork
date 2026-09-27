"""Re-split the AA cleaning EXPENSES between the HOA and the owners, per the workbook.

`fixes/yacinde_hoa_split` moves owner-borne lines ONTO the receivable, reads `Invoice - HOA`
(the HOA's half only) and queries **Bill** objects.  None of that fits a correction: cleaning
now arrives as an Expense (owner decision 2026-09-19), and when the owner revises the
allocation a clean can move in EITHER direction.

So this reads `Cleaning Allocation`, which is the only tab carrying `hoa_paid` and
`owner_paid` side by side, and repoints **Purchase** lines both ways:

    hoa_paid  > 0   ->  Trust Assets other than Cash & A/R:HOA Receivable - Yacinde
    owner_paid > 0  ->  1C - Owner Expenses:Cleaning Expense - Owner

Only `AccountRef` moves.  The class stays exactly as it is: both halves already carry the
LISTING class (owner decision 2026-09-17), so the class is not what distinguishes them and
rewriting it would be a second, invisible change riding along with this one.

This is the other half of `invoices/hoa_cleaning_update`, and neither is correct alone.
Lowering the HOA's invoice without this leaves the cost on the receivable, which then reads
as "laid out and not yet billed" when it is really "not the HOA's at all".  Doing this
without the invoice leaves the HOA billed for a clean it no longer bears.

Which expense is touched is never guessed: the workbook's `invoice` column names the AA
invoice, and the DocNumber scheme is `<yyyymmdd>_INV<aa invoice>_<total>`.  A row whose
AA invoice has no Expense on the books is reported, never silently dropped.

    python -m src.fixes.yacinde_hoa_resplit --period 202607            # dry run + CSV
    python -m src.fixes.yacinde_hoa_resplit                            # every workbook
    python -m src.fixes.yacinde_hoa_resplit --period 202607 --confirm  # write
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

from ..config import client, name as acct_name_of
from ..paths import review_csv
from ..qbo_client import QBOClient
from ..resolver import Resolver
from ..invoices.hoa_cleaning import WORKBOOK_DIR, workbooks

ALLOC_SHEET = "Cleaning Allocation"
VENDOR = "AA Professional Cleaners LLC"
# `07/13/2026 C1 Cleaning Fee` -- and a description can contain a NEWLINE, so re.S.
DESC_RE = re.compile(r"^\s*(\d{2})/(\d{2})/(\d{4})\s+(\S+)\s+(.*)$", re.S)


def iso(v) -> str:
    return str(v)[:10]


def read_alloc(path: Path) -> list[dict]:
    """One row per clean that COST money, off `Cleaning Allocation`.

    `$0` rows are fraction-week "no clean" markers, not cleans (CLAUDE.md), and a row with
    no `invoice` was never billed by AA, so neither reaches an expense line.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if ALLOC_SHEET not in wb.sheetnames:
        sys.exit(f"{path.name}: no {ALLOC_SHEET!r} sheet")
    rows = wb[ALLOC_SHEET].iter_rows(values_only=True)
    hdr = list(next(rows))
    out = []
    for r in rows:
        d = dict(zip(hdr, r))
        unit = str(d.get("unit") or "").strip()
        if not unit or unit.upper() == "TOTAL" or d.get("date") is None:
            continue
        amt = round(float(d.get("Amount") or 0), 2)
        inv = str(d.get("invoice") or "").strip()
        if amt <= 0 or not inv or not inv.isdigit():
            continue
        hoa = round(float(d.get("hoa_paid") or 0), 2)
        own = round(float(d.get("owner_paid") or 0), 2)
        if round(hoa + own, 2) != amt:
            sys.exit(f"{path.name}: {iso(d['date'])} {unit} ${amt:,.2f} — "
                     f"hoa {hoa:,.2f} + owner {own:,.2f} does not equal it")
        out.append({"Date": iso(d["date"]), "Unit": unit, "Amount": amt, "AAInvoice": inv,
                    "Side": "HOA" if hoa > 0 else "OWNER",
                    "Party": str(d.get("party") or "").strip()})
    return out


def line_key(line: dict) -> tuple[str | None, str, float]:
    det = line.get("AccountBasedExpenseLineDetail") or {}
    m = DESC_RE.match(line.get("Description") or "")
    date = f"{m.group(3)}-{m.group(1)}-{m.group(2)}" if m else None
    leaf = ((det.get("ClassRef") or {}).get("name") or "").rsplit(":", 1)[-1].strip()
    unit = leaf or (f"Yacinde {m.group(4)}" if m else "")
    return date, unit, round(float(line.get("Amount") or 0), 2)


def match_lines(alloc: list[dict], lines: list[dict]) -> tuple[dict[int, dict], list[str]]:
    """line index -> allocation row.  Exact (date, unit, amount) first; then unit relaxed;
    then date relaxed, LAST, so it can never steal a clean from a better pair."""
    pool: dict[tuple, list[dict]] = defaultdict(list)
    for a in alloc:
        pool[(a["Date"], a["Unit"], a["Amount"])].append(a)
    by_du: dict[tuple, list[dict]] = defaultdict(list)
    by_ua: dict[tuple, list[dict]] = defaultdict(list)
    for a in alloc:
        by_du[(a["Date"], a["Amount"])].append(a)
        by_ua[(a["Unit"], a["Amount"])].append(a)

    used, out, notes = set(), {}, []

    def take(bucket, key):
        for a in bucket.get(key, []):
            if id(a) not in used:
                used.add(id(a))
                return a
        return None

    todo = list(enumerate(lines))
    for pas, (bucket, keyf, label) in enumerate((
            (pool, lambda k: k, None),
            (by_du, lambda k: (k[0], k[2]), "UNIT relaxed"),
            (by_ua, lambda k: (k[1], k[2]), "DATE relaxed"))):
        nxt = []
        for i, ln in todo:
            k = line_key(ln)
            if k[0] is None and pas < 2:
                nxt.append((i, ln))
                continue
            a = take(bucket, keyf(k))
            if a is None:
                nxt.append((i, ln))
                continue
            out[i] = a
            if label:
                notes.append(f"{label}: line {k} paired with "
                             f"({a['Date']}, {a['Unit']}, {a['Amount']:.2f})")
        todo = nxt
    for i, ln in todo:
        notes.append(f"UNMATCHED line: {line_key(ln)}  {(ln.get('Description') or '')[:60]!r}")
    leftover = [a for a in alloc if id(a) not in used]
    for a in leftover:
        notes.append(f"UNMATCHED workbook row: ({a['Date']}, {a['Unit']}, {a['Amount']:.2f}) "
                     f"AA {a['AAInvoice']}")
    return out, notes


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--period", default=None, help="YYYYMM; default = every workbook found")
    ap.add_argument("--out", default=None)
    ap.add_argument("--confirm", action="store_true", help="actually WRITE to QuickBooks")
    args = ap.parse_args()

    books = workbooks(args.period)
    if not books:
        sys.exit(f"no workbooks matching period {args.period!r} in {WORKBOOK_DIR}")

    qbo = client()
    res = Resolver(qbo)
    hoa_name = acct_name_of("hoa_receivable_yacinde")
    own_name = acct_name_of("cleaning_expense_owner")
    target = {"HOA": res.account(hoa_name), "OWNER": res.account(own_name)}
    print("resolving:")
    for side, nm in (("HOA", hoa_name), ("OWNER", own_name)):
        print(f"  {side:<6} {nm[:66]:<66} "
              f"{'Id ' + target[side] if target[side] else '*** MISSING ***'}")
    if not all(target.values()):
        sys.exit("ABORT — an account does not exist")

    alloc: list[dict] = []
    for b in books:
        got = read_alloc(b)
        print(f"\n{b.name}: {len(got)} charged clean(s), "
              f"HOA ${sum(a['Amount'] for a in got if a['Side']=='HOA'):,.2f} / "
              f"owner ${sum(a['Amount'] for a in got if a['Side']=='OWNER'):,.2f}, "
              f"AA invoice(s) {sorted({a['AAInvoice'] for a in got})}")
        alloc.extend(got)

    # `EntityRef` is not a queryable property on Purchase, so the vendor cannot be the
    # filter.  The DocNumber scheme `<yyyymmdd>_INV<aa invoice>_<total>` is, and it is the
    # same string the workbook's `invoice` column names.
    wanted = sorted({a["AAInvoice"] for a in alloc})
    purchases: dict[str, dict] = {}
    by_inv: dict[str, dict] = {}
    for num in wanted:
        got = qbo.query_all(
            f"SELECT * FROM Purchase WHERE DocNumber LIKE '%INV{num}\\_%'", "Purchase")
        got = [g for g in got
               if re.search(rf"_INV{num}_", g.get("DocNumber") or "")
               and (g.get("EntityRef") or {}).get("name") == VENDOR]
        if len(got) > 1:
            sys.exit(f"ABORT — {len(got)} Expenses for AA invoice {num}: "
                     f"{[g['DocNumber'] for g in got]}")
        if got:
            purchases[got[0]["Id"]] = got[0]
            by_inv[num] = got[0]

    rows, edits_by_p, all_notes = [], defaultdict(list), []
    missing = [n for n in wanted if n not in by_inv]
    if missing:
        print(f"\n*** no Expense on the books for AA invoice(s) {missing} ***")

    for num in wanted:
        p = by_inv.get(num)
        if p is None:
            continue
        mine = [a for a in alloc if a["AAInvoice"] == num]
        lines = [l for l in p.get("Line", []) if l.get("AccountBasedExpenseLineDetail")]
        pairs, notes = match_lines(mine, lines)
        all_notes += [f"{p['DocNumber']}: {n}" for n in notes]

        wb_tot = round(sum(a["Amount"] for a in mine), 2)
        ln_tot = round(sum(float(l.get("Amount") or 0) for l in lines), 2)
        print(f"\n{'='*78}\nAA {num}  ->  Expense {p['Id']}  {p['DocNumber']}  {p['TxnDate']}"
              f"  Total {float(p['TotalAmt']):,.2f}")
        print(f"  workbook {len(mine):>3} cleans ${wb_tot:>10,.2f}"
              f"   |  expense {len(lines):>3} lines ${ln_tot:>10,.2f}"
              + ("   *** they do not agree ***" if abs(wb_tot - ln_tot) > 0.005 else "   (tie)"))

        before = defaultdict(float)
        after = defaultdict(float)
        for i, l in enumerate(lines):
            det = l["AccountBasedExpenseLineDetail"]
            cur = det.get("AccountRef", {}).get("value")
            cur_nm = det.get("AccountRef", {}).get("name", "")
            before[cur_nm] += float(l.get("Amount") or 0)
            a = pairs.get(i)
            if a is None:
                after[cur_nm] += float(l.get("Amount") or 0)
                continue
            want = target[a["Side"]]
            want_nm = hoa_name if a["Side"] == "HOA" else own_name
            after[want_nm] += float(l.get("Amount") or 0)
            if cur != want:
                edits_by_p[p["Id"]].append((i, want, a))
            rows.append({
                "DocNumber": p["DocNumber"], "PurchaseId": p["Id"], "TxnDate": p["TxnDate"],
                "LineIdx": i, "CleanDate": a["Date"], "Unit": a["Unit"],
                "Amount": f"{float(l.get('Amount') or 0):.2f}", "AAInvoice": num,
                "Class": (det.get("ClassRef") or {}).get("name", ""),
                "FromAccount": cur_nm, "ToAccount": want_nm, "Side": a["Side"],
                "Party": a["Party"],
                "Action": "MOVE" if cur != want else "already correct",
                "Description": (l.get("Description") or "").replace("\n", " ")})
        for nm in sorted(set(before) | set(after)):
            d = after[nm] - before[nm]
            print(f"    {nm[-52:]:<52} {before[nm]:>10,.2f} -> {after[nm]:>10,.2f}  {d:+10,.2f}")
        n_move = len(edits_by_p.get(p["Id"], []))
        print(f"    {n_move} line(s) move" + ("" if n_move else "  (already correct)"))

    if all_notes:
        print(f"\n{'-'*78}\nflags ({len(all_notes)}):")
        for n in all_notes:
            print(f"  ! {n}")

    if not rows:
        sys.exit("\nnothing matched — nothing to do")

    out = Path(args.out) if args.out else review_csv("yacinde_hoa_resplit")
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)

    moved = sum(len(v) for v in edits_by_p.values())
    to_hoa = sum(float(r["Amount"]) for r in rows if r["Action"] == "MOVE" and r["Side"] == "HOA")
    to_own = sum(float(r["Amount"]) for r in rows if r["Action"] == "MOVE" and r["Side"] == "OWNER")
    print(f"\n{'='*78}\n{len(rows)} line(s) over {len(edits_by_p)} expense(s) -> {out}")
    print(f"  {moved} move: ${to_own:,.2f} onto the owners, ${to_hoa:,.2f} onto the receivable")
    print(f"  net effect on {hoa_name.rsplit(':',1)[-1]}: {to_hoa - to_own:+,.2f}")

    if not args.confirm:
        print("\nDRY RUN — nothing written. Review the CSV, then re-run with --confirm.")
        return

    for pid, edits in edits_by_p.items():
        p = purchases[pid]
        body = json.loads(json.dumps(p))
        lines = [l for l in body.get("Line", []) if l.get("AccountBasedExpenseLineDetail")]
        for i, want, _a in edits:
            lines[i]["AccountBasedExpenseLineDetail"]["AccountRef"] = {"value": want}
        body["sparse"] = False
        before_tot = round(float(p["TotalAmt"]), 2)
        got = qbo.post("purchase", body)
        after_tot = round(float(got.get("TotalAmt") or 0), 2)
        if abs(after_tot - before_tot) > 0.005:
            sys.exit(f"ABORT — {p['DocNumber']}: total changed "
                     f"{before_tot:,.2f} -> {after_tot:,.2f}")
        print(f"  Expense {got['Id']}  {got.get('DocNumber')}  ${after_tot:,.2f}  "
              f"{len(edits)} line(s) repointed  SyncToken {got['SyncToken']}")
    print(f"\nrepointed {moved} line(s) over {len(edits_by_p)} expense(s)")


if __name__ == "__main__":
    main()
