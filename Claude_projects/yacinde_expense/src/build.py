"""Combine Yacinde bookings, label ownership, match cleans, and allocate cleaning + supplies.

    python -m src.fetch_guesty          # refresh data/guesty/guesty_reservations.json first
    python -m src.build [--start YYYY-MM-DD --end YYYY-MM-DD]

Steps
  1. Load previous-company arrivals (Hold blocks dropped) and Guesty reservations
     (confirmed + canceled, including owner / owner-guest stays).
  2. Dedupe: exact unit+check-in+check-out match merges the two sources. Before the
     migration cutoff the previous-company status wins; on/after it Guesty wins.
     Previous-company rows on/after the cutoff that are not in Guesty but whose guest
     surname appears in Guesty for the same unit are treated as re-dated in Guesty.
  3. Label each booking with the fractional-master week owner (week containing check-in).
  4. Supplies = guests x nights x rate; timeshare (Individual-owner) stays use max sleeps.
  5. Match each clean to the latest active checkout in that unit within N days before it.
Writes output/yacinde_expense_allocation_<start>_<end>.xlsx.
"""
import argparse
import datetime
import json
import re
from pathlib import Path

import pandas as pd
import yaml
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
TZ = "America/Los_Angeles"
NOT_IN_MASTER = "Not in fractional master"


def load_cfg() -> dict:
    return yaml.safe_load((ROOT / "config" / "allocation_rules.yml").read_text())


def _surname(name: str) -> str:
    name = re.sub(r"\(.*?\)", "", str(name or "")).strip()
    return name.split()[-1].lower() if name.split() else ""


# ---------------------------------------------------------------- loaders
def load_guesty(path: Path) -> pd.DataFrame:
    rows = []
    for x in json.loads(path.read_text()):
        lst = x.get("listing") or {}
        guest = (x.get("guest") or {}).get("fullName") or ""
        rows.append({
            "unit": lst.get("nickname"),
            "check_in": pd.Timestamp(x["checkIn"]).tz_convert(TZ).tz_localize(None).normalize(),
            "check_out": pd.Timestamp(x["checkOut"]).tz_convert(TZ).tz_localize(None).normalize(),
            "status": x["status"],
            "channel": x.get("source"),
            "guest": guest.strip(),
            "guesty_code": x.get("confirmationCode"),
            "guesty_guests": x.get("guestsCount"),
            "max_sleeps": lst.get("accommodates"),
            # owner reservations migrated from the previous company carry guestsCount = accommodates
            "guests_placeholder": x.get("source") == "owner" and "(NWV)" in guest,
        })
    return pd.DataFrame(rows)


def load_previous(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    p = pd.read_excel(path)
    holds = p[p["Res. Status"] == "Hold"]
    p = p[p["Res. Status"] != "Hold"]
    df = pd.DataFrame({
        "unit": p["Unit Name"],
        "check_in": pd.to_datetime(p["Check-in Date"]).dt.normalize(),
        "check_out": pd.to_datetime(p["Checkout Date"]).dt.normalize(),
        "status": "confirmed",
        "channel": p["Res Type"],
        "guest": (p["First Name"].fillna("").astype(str) + " " + p["Last Name"].fillna("").astype(str)).str.strip(),
        "prev_res_no": p["Res. #"].astype(str),
        "prev_alt_confirm": p["Alt. Confirm #"],
    })
    return df, holds


def load_master(path: Path, sheet: str) -> pd.DataFrame:
    m = pd.read_excel(path, sheet_name=sheet)
    m["checkin"] = pd.to_datetime(m["checkin"]).dt.normalize()
    m["checkout"] = pd.to_datetime(m["checkout"]).dt.normalize()
    # B1 segment K is listed twice (Nugrowth + William & Loreta Lynn): the individual owner wins.
    m["_rank"] = (m["Ownership.Type"] != "Individual").astype(int)
    return m.sort_values("_rank").drop_duplicates(["Listing", "checkin"]).drop(columns="_rank")


def load_whole_owners(path: Path) -> dict:
    """Unit -> owner for whole-owned units (UNIT TITLE without a fractional segment letter, e.g. 'B6')."""
    o = pd.read_excel(path)
    o = o[o["UNIT TITLE"].astype(str).str.fullmatch(r"[A-Z]\d")]
    return {"Yacinde " + t: str(n).strip().rstrip(".").replace("NuGrowth", "Nugrowth")   # match fractional master spelling
            for t, n in zip(o["UNIT TITLE"], o["OWNER NAME"])}


def load_cleaning(path: Path) -> pd.DataFrame:
    c = pd.read_excel(path, sheet_name="Cleaning Records")
    c["Date"] = pd.to_datetime(c["Date"]).dt.normalize()
    c["unit"] = "Yacinde " + c["Unit"].astype(str).str.strip()
    c["Invoice #"] = c["Invoice #"].astype(str)
    return c.reset_index(names="clean_id")


# ---------------------------------------------------------------- combine
def combine(prev: pd.DataFrame, gy: pd.DataFrame, cutoff: pd.Timestamp, cleans: pd.DataFrame, max_days: int):
    log = []
    key = ["unit", "check_in", "check_out"]

    # Guesty internal duplicates (e.g. a canceled manual copy of a stay): keep confirmed first
    gy = gy.assign(_r=(gy.status != "confirmed").astype(int)).sort_values(key + ["_r"])
    for _, r in gy[gy.duplicated(key, keep="first")].iterrows():
        log.append({**r[key + ["guest", "status"]].to_dict(), "source": "guesty", "ref": r.guesty_code,
                    "action": "dropped: duplicate Guesty row for same unit/dates"})
    gy = gy.drop_duplicates(key, keep="first").drop(columns="_r")

    m = prev.merge(gy, on=key, how="outer", suffixes=("_prev", "_gy"), indicator=True)
    out = []
    for _, r in m.iterrows():
        both, pre = r["_merge"] == "both", r.check_in < cutoff
        if r["_merge"] == "left_only" and not pre:
            same_guest = gy[(gy.unit == r.unit) & (gy.guest.map(_surname) == _surname(r.guest_prev))]
            in_gy = ", ".join(f"{g.guesty_code} {g.check_in:%m/%d}-{g.check_out:%m/%d} ({g.status})"
                              for g in same_guest.itertuples())
            overlaps = ((same_guest.check_in < r.check_out) & (same_guest.check_out > r.check_in)).any()
            cleaned = cleans[(cleans.unit == r.unit) & cleans.Date.between(
                r.check_out, r.check_out + pd.Timedelta(days=max_days))]
            base = {**r[key].to_dict(), "guest": r.guest_prev, "status": r.status_prev,
                    "source": "previous_company", "ref": r.prev_res_no}
            if len(same_guest) and not overlaps and len(cleaned):
                # Guesty holds the same guest on different, non-overlapping dates, yet a clean
                # follows this checkout: the original stay happened, keep it.
                log.append({**base, "action": f"KEPT: clean on {cleaned.Date.min():%m/%d} supports this stay; "
                                              f"Guesty also has {in_gy} (possible re-date - verify)"})
            else:
                log.append({**base, "action": f"dropped: re-dated in Guesty as {in_gy}" if len(same_guest)
                                              else "dropped: on/after migration cutoff but not in Guesty"})
                continue
        if both:
            status = r.status_prev if pre else r.status_gy
            if r.status_prev != r.status_gy:
                log.append({**r[key].to_dict(), "guest": r.guest_gy, "status": f"prev={r.status_prev} / guesty={r.status_gy}",
                            "source": "both", "ref": f"{r.prev_res_no} / {r.guesty_code}",
                            "action": f"status conflict: used {'previous company' if pre else 'Guesty'} ({status})"})
        out.append({
            "unit": r.unit, "check_in": r.check_in, "check_out": r.check_out,
            "source": "both" if both else ("previous_company" if r["_merge"] == "left_only" else "guesty"),
            "status": status if both else (r.status_prev if r["_merge"] == "left_only" else r.status_gy),
            "channel": r.channel_gy if pd.notna(r.channel_gy) else r.channel_prev,
            "prev_channel": r.channel_prev,
            "guest": r.guest_gy if isinstance(r.guest_gy, str) and r.guest_gy else r.guest_prev,
            "guesty_code": r.guesty_code, "prev_res_no": r.prev_res_no, "prev_alt_confirm": r.prev_alt_confirm,
            "guesty_guests": r.guesty_guests, "guests_placeholder": bool(r.guests_placeholder) if both or r["_merge"] == "right_only" else False,
        })
    b = pd.DataFrame(out)
    b["nights"] = (b.check_out - b.check_in).dt.days
    return b.sort_values(["check_in", "unit"]).reset_index(drop=True), pd.DataFrame(log)


def add_owner_weeks(b: pd.DataFrame, master: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Add Individual (timeshare) owner weeks from the fractional master that neither booking source
    has, for weeks starting on/after `start` (first week 06/26-07/03, before either source's coverage).
    A week is skipped if any booking (even a canceled one) overlaps it in that unit."""
    w = master[(master["Ownership.Type"] == "Individual") & (master.checkin >= start) & (master.checkin <= end)]
    add = []
    for r in w.itertuples():
        overlap = b[(b.unit == r.Listing) & (b.check_in < r.checkout) & (b.check_out > r.checkin)]
        if len(overlap):
            continue
        add.append({"unit": r.Listing, "check_in": r.checkin, "check_out": r.checkout, "source": "fractional_master",
                    "status": "confirmed", "channel": "owner week (fractional master)", "prev_channel": None,
                    "guest": r.Owner, "guesty_code": None, "prev_res_no": None, "prev_alt_confirm": None,
                    "guesty_guests": None, "guests_placeholder": False,
                    "nights": (r.checkout - r.checkin).days})
    added = pd.DataFrame(add)
    if len(added):
        b = pd.concat([b, added], ignore_index=True).sort_values(["check_in", "unit"]).reset_index(drop=True)
    return b, added


# ---------------------------------------------------------------- ownership + supplies
def label_ownership(b: pd.DataFrame, master: pd.DataFrame, whole: dict, parties: dict) -> pd.DataFrame:
    cols = {"Owner": "owner", "Ownership.Type": "ownership_type", "Segment": "segment",
            "Week.#": "week_no", "checkin": "week_start", "checkout": "week_end",
            "In.Rental.Program": "in_rental_program"}
    labelled = []
    for r in b.itertuples():
        wk = master[(master.Listing == r.unit) & (master.checkin <= r.check_in) & (master.checkout > r.check_in)]
        info = {v: None for v in cols.values()}
        if len(wk):
            w = wk.iloc[0]
            info = {v: w[k] for k, v in cols.items()}
            nxt = master[(master.Listing == r.unit) & (master.checkin < r.check_out) & (master.checkout > r.check_in)]
            info["spans_weeks"] = len(nxt) > 1
            info["other_week_owners"] = "; ".join(sorted(set(nxt.Owner) - {w.Owner})) or None
        elif r.unit in whole:
            info["owner"], info["ownership_type"] = whole[r.unit], "Whole owner"
            info["spans_weeks"], info["other_week_owners"] = None, None
        else:
            info["ownership_type"] = NOT_IN_MASTER
            info["spans_weeks"], info["other_week_owners"] = None, None
        labelled.append(info)
    b = pd.concat([b, pd.DataFrame(labelled, index=b.index)], axis=1)
    b["party"] = b.ownership_type.map(parties).fillna("UNASSIGNED")
    b["is_timeshare"] = b.ownership_type == "Individual"
    return b


def add_supplies(b: pd.DataFrame, max_sleeps: dict, rate: float) -> pd.DataFrame:
    b["max_sleeps"] = b.unit.map(max_sleeps)
    guests, basis = [], []
    for r in b.itertuples():
        if r.is_timeshare:
            guests.append(r.max_sleeps); basis.append("max sleeps (timeshare)")
        elif r.source in ("previous_company", "fractional_master"):
            guests.append(r.max_sleeps); basis.append("max sleeps (no guest count in previous-company file)")
        elif r.guests_placeholder:
            guests.append(r.max_sleeps); basis.append("max sleeps (migrated owner reservation, no real count)")
        else:
            guests.append(r.guesty_guests); basis.append("Guesty guestsCount")
    b["supply_guests"], b["supply_guests_basis"] = guests, basis
    b["active"] = b.status == "confirmed"
    b["supplies_cost"] = (b.supply_guests * b.nights * rate).where(b.active, 0.0).round(2)
    return b


# ---------------------------------------------------------------- cleaning
def match_cleans(c: pd.DataFrame, b: pd.DataFrame, max_days: int, whole: dict, parties: dict) -> pd.DataFrame:
    """Merge each clean to a checkout in the same unit: on the checkout date, or up to `max_days` later.

    Pass 1 matches clean date == checkout date. Pass 2 matches remaining cleans to the most recent
    still-uncleaned checkout 1..max_days days earlier, but not when another guest is mid-stay on the
    clean date. One clean per stay; anything else is left unmatched (still allocated via the unit's
    fraction week / whole owner).
    """
    act = b[b.active]
    used, hits = set(), {}
    for exact in (True, False):
        for r in c.sort_values(["Date", "clean_id"]).itertuples():
            if r.clean_id in hits:
                continue
            u = act[act.unit == r.unit]
            if not exact and ((u.check_in < r.Date) & (u.check_out > r.Date)).any():
                continue                                  # a guest is in the unit: not a late turnover
            gap = (r.Date - u.check_out).dt.days
            cand = u[(gap == 0) if exact else gap.between(1, max_days)]
            cand = cand[~cand.index.isin(used)]
            if len(cand):
                hit = cand.sort_values("check_out").iloc[-1]
                used.add(hit.name)
                g = (r.Date - hit.check_out).days
                hits[r.clean_id] = (hit.name, "checkout date" if g == 0 else f"{g} day(s) after checkout")
    rows = []
    for r in c.itertuples():
        if r.clean_id in hits:
            idx, how = hits[r.clean_id]
            rows.append({"clean_id": r.clean_id, "match": how, "match_note": None, "booking_idx": idx})
            continue
        near = b[(b.unit == r.unit) & (b.check_out <= r.Date) & (b.check_out >= r.Date - pd.Timedelta(days=max_days))]
        note = "; ".join(f"{x.guest} out {x.check_out:%m/%d} ({x.status}{', already cleaned' if x.Index in used else ''})"
                         for x in near.itertuples()) or None
        rows.append({"clean_id": r.clean_id, "match": "UNMATCHED", "match_note": note, "booking_idx": None})
    mm = pd.DataFrame(rows)
    bcols = ["unit", "check_in", "check_out", "nights", "guest", "channel", "source", "guesty_code",
             "prev_res_no", "week_start", "week_end", "owner", "ownership_type", "party"]
    bb = b[bcols].drop(columns="unit").add_prefix("bk_")
    out = c.merge(mm, on="clean_id").merge(bb, left_on="booking_idx", right_index=True, how="left")
    out["party"] = out.bk_party.fillna("UNALLOCATED (no matching stay)")
    out["ownership_type"] = out.bk_ownership_type
    out["owner"] = out.bk_owner
    # whole-owned units have one owner all year, so a clean with no matching stay still has an owner
    nostay = out.bk_party.isna() & out.unit.isin(whole)
    out.loc[nostay, "owner"] = out.loc[nostay, "unit"].map(whole)
    out.loc[nostay, "ownership_type"] = "Whole owner"
    out.loc[nostay, "party"] = parties["Whole owner"]
    return out.drop(columns=["bk_party", "bk_ownership_type", "bk_owner"])


def fraction_weeks(cl: pd.DataFrame, master: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp,
                   parties: dict) -> pd.DataFrame:
    """Tag each clean on a fractional unit with the fraction week it falls in, and record a $0 HOA
    clean for every fraction week in the period that has no invoiced clean.

    A clean belongs to the week with week_start < clean date <= week_end, so a Friday checkout clean
    counts for the week that just ended. That week's owner pays any clean HOA does not cover.
    Weeks count toward the period when week_end is within [start, end] (first: 06/26-07/03).
    """
    cl = cl.copy()
    cols = {"checkin": "week_start", "checkout": "week_end", "Segment": "segment", "Owner": "owner",
            "Ownership.Type": "ownership_type"}
    for c in cols.values():
        if c not in cl:
            cl[c] = None
    fu = set(master.Listing)
    for i, r in cl[cl.unit.isin(fu)].iterrows():
        wk = master[(master.Listing == r.unit) & (master.checkin < r.Date) & (master.checkout >= r.Date)]
        if len(wk):
            for k, v in cols.items():
                cl.at[i, v] = wk.iloc[0][k]
            cl.at[i, "party"] = parties.get(cl.at[i, "ownership_type"], "UNASSIGNED")

    weeks = master[master.checkout.between(start, end)]
    have = set(zip(cl.unit, pd.to_datetime(cl.week_start)))
    miss = weeks[[(u, w) not in have for u, w in zip(weeks.Listing, weeks.checkin)]]
    if len(miss):
        synth = pd.DataFrame({
            "clean_id": range(int(cl.clean_id.max()) + 1, int(cl.clean_id.max()) + 1 + len(miss)),
            "Date": miss.checkout.values, "Unit": miss.Listing.str.replace("Yacinde ", "").values,
            "Rate": 0, "Qty": 1, "Amount": 0.0, "Invoice #": "not invoiced", "Invoice date": pd.NaT,
            "Billing period": miss.checkout.dt.strftime("%b %Y").values,
            "Notes": "Fraction week with no invoiced clean - HOA clean recorded at $0",
            "unit": miss.Listing.values, "match": "fraction week (no clean)",
            "week_start": miss.checkin.values, "week_end": miss.checkout.values, "segment": miss.Segment.values,
            "owner": miss.Owner.values, "ownership_type": miss["Ownership.Type"].values,
            "party": miss["Ownership.Type"].map(parties).values})
        cl = pd.concat([cl, synth], ignore_index=True)
    return cl


def assign_cleaning_paid(cl: pd.DataFrame, hoa: dict) -> pd.DataFrame:
    """Add cleaning_paid_by / hoa_paid / owner_paid / hoa_rule.

    * Whole-owner units: HOA pays the first N eligible cleans per calendar month (per unit).
    * Fractional units: every fraction week (any owner) gets N HOA cleans (the first N by date);
      further cleans that week are paid by the week's owner.
    * Everything else is paid by the allocated party.
    """
    cl = cl.copy()
    cl["month"] = cl.Date.dt.strftime("%Y-%m")
    excluded = cl["Invoice #"].isin([str(i) for i in hoa["exclude_invoices"]])
    order = ["Date", "clean_id"]

    whole = cl[(cl.ownership_type == "Whole owner") & ~excluded].sort_values(order)
    group = ["month", "unit"] if hoa["scope"] == "unit" else ["month"]
    seq_whole = whole.groupby(group).cumcount() + 1

    ts = cl[cl.week_start.notna() & ~excluded].sort_values(order)
    seq_ts = ts.groupby(["unit", "week_start"]).cumcount() + 1

    cl["hoa_clean_seq"] = pd.concat([seq_whole, seq_ts])        # NaN for cleans outside both rules
    cl["hoa_rule"] = None
    cl.loc[whole.index, "hoa_rule"] = f"whole owner: first {hoa['free_cleans_per_month']}/month"
    cl.loc[ts.index, "hoa_rule"] = f"fraction: first {hoa['fraction_free_cleans_per_week']}/fraction week"
    hoa_paid = (cl.index.isin(whole.index) & (cl.hoa_clean_seq <= hoa["free_cleans_per_month"])) | \
               (cl.index.isin(ts.index) & (cl.hoa_clean_seq <= hoa["fraction_free_cleans_per_week"]))
    cl["cleaning_paid_by"] = cl.party.where(~hoa_paid, hoa["paid_by"])
    cl["hoa_paid"] = cl.Amount.where(hoa_paid, 0.0)
    cl["owner_paid"] = cl.Amount - cl.hoa_paid
    return cl


def bookings_with_cleans(bk: pd.DataFrame, cl: pd.DataFrame, master: pd.DataFrame,
                         start: pd.Timestamp, end: pd.Timestamp, parties: dict) -> pd.DataFrame:
    """One sheet built on the owner fraction weeks: each week, the stays checking in during it,
    and each stay's clean(s) merged on its checkout.

    Fractional units (B1 B2 B3 B4 F1) get a block per master week (week_end in the period, or any
    week a stay/clean falls in); a week with no stay still gets a row, carrying its $0 HOA clean if any.
    Cleans with no stay sit in the week they are allocated to, except that one dated on a checkout
    with no clean of its own (e.g. the week-end $0 HOA clean) shares that stay's row. Whole-owner units have no weeks, so
    their stays are listed by date. A stay with two cleans takes two rows; supplies are shown on the
    first row only so the column sums.
    """
    wk_cols = ["week_start", "week_end", "segment", "owner", "ownership_type", "party"]
    stay_cols = ["check_in", "check_out", "nights", "guest", "channel", "source", "supply_guests",
                 "supply_guests_basis", "supplies_cost", "guesty_code", "prev_res_no"]
    clean_cols = ["Date", "Amount", "Invoice #", "match", "cleaning_paid_by", "hoa_paid", "owner_paid", "hoa_rule",
                  "hoa_clean_seq", "owner", "ownership_type", "week_start", "week_end", "Notes"]
    st = bk[["unit"] + wk_cols].join(bk[stay_cols].add_prefix("stay_"))
    st["_bidx"] = bk.index.astype(float)
    cc = cl[clean_cols].add_prefix("clean_")
    cc["_bidx"] = pd.to_numeric(cl.booking_idx, errors="coerce")
    linked = cc._bidx.isin(st._bidx)
    for i in cl.index[~linked]:                   # same-day checkout with no clean of its own
        free = st[(st.unit == cl.at[i, "unit"]) & (st.stay_check_out == cl.at[i, "Date"]) & ~st._bidx.isin(cc._bidx)]
        if len(free):
            cc.at[i, "_bidx"] = free._bidx.iloc[0]
            linked[i] = True
    loose = cl.loc[~linked, ["unit"] + wk_cols].join(cc[~linked])      # cleans with no stay: in their own week

    m = master.rename(columns={"Listing": "unit", "checkin": "week_start", "checkout": "week_end",
                               "Segment": "segment", "Owner": "owner", "Ownership.Type": "ownership_type"})
    m["party"] = m.ownership_type.map(parties).fillna("UNASSIGNED")
    keys = set(zip(st.unit, st.week_start)) | set(zip(loose.unit, loose.week_start))
    in_scope = m.week_end.between(start, end) | pd.Series([k in keys for k in zip(m.unit, m.week_start)], index=m.index)
    weeks = m.loc[in_scope, ["unit"] + wk_cols]
    filled = set(zip(st.unit, st.week_start)) | set(zip(loose.unit, loose.week_start))
    empty = weeks[[k not in filled for k in zip(weeks.unit, weeks.week_start)]]
    out = pd.concat([st.merge(cc[linked], on="_bidx", how="left"), loose, empty], ignore_index=True)
    out["_date"] = out.stay_check_in.fillna(out.clean_Date).fillna(out.week_end)
    out["_week"] = out.week_start.fillna(out._date)
    out = out.sort_values(["unit", "_week", "_date", "stay_check_in", "clean_Date"], na_position="last")
    out.loc[out._bidx.notna() & out.duplicated("_bidx"), "stay_supplies_cost"] = None
    return out.drop(columns=["_bidx", "_date", "_week"]).reset_index(drop=True)


# ---------------------------------------------------------------- main
def main():
    cfg = load_cfg()
    ap = argparse.ArgumentParser(description="Yacinde cleaning + supplies allocation")
    ap.add_argument("--start", default=str(cfg["period"]["start"]))
    ap.add_argument("--end", default=str(cfg["period"]["end"]))
    args = ap.parse_args()
    start, end = pd.Timestamp(args.start), pd.Timestamp(args.end)
    inp = {k: ROOT / v if k != "fractional_sheet" else v for k, v in cfg["inputs"].items()}

    gy = load_guesty(inp["guesty_json"])
    prev, holds = load_previous(inp["previous_company"])
    master = load_master(inp["fractional_master"], inp["fractional_sheet"])
    whole = {u: o for u, o in load_whole_owners(inp["whole_owners"]).items() if u not in set(master.Listing)}
    clean = load_cleaning(inp["cleaning"])

    max_days = cfg["cleaning_match"]["max_days_after_checkout"]
    b, dedup_log = combine(prev, gy.drop(columns="max_sleeps"), pd.Timestamp(str(cfg["migration_cutoff"])), clean, max_days)
    b, owner_weeks = add_owner_weeks(b, master, pd.Timestamp(str(cfg["timeshare_weeks_from"])), end)
    b = label_ownership(b, master, whole, cfg["parties"])
    b = add_supplies(b, gy.groupby("unit").max_sleeps.max().to_dict(), cfg["supplies"]["rate_per_guest_night"])
    b["in_period"] = b.check_out.between(start, end)
    cl = match_cleans(clean, b, max_days, whole, cfg["parties"])
    cl = fraction_weeks(cl, master, start, end, cfg["parties"])
    cl = assign_cleaning_paid(cl.sort_values(["Date", "unit", "clean_id"]).reset_index(drop=True), cfg["hoa_cleaning"])

    # ---- summaries (stays counted by check-out within the period)
    bp = b[b.in_period & b.active]
    stays = bp.groupby(["party", "ownership_type"]).agg(stays=("unit", "size"), nights=("nights", "sum"),
                                                        supplies_cost=("supplies_cost", "sum"))
    cleans = cl.groupby([cl.cleaning_paid_by, cl.ownership_type.fillna("-")]).agg(
        cleans=("Amount", "size"), cleaning_cost=("Amount", "sum"))
    cleans.index.names = ["party", "ownership_type"]
    summary = stays.join(cleans, how="outer").fillna(0).reset_index()
    summary["total_cost"] = summary.supplies_cost + summary.cleaning_cost
    summary.loc[len(summary)] = ["TOTAL", "", *summary.iloc[:, 2:].sum().tolist()]

    by_owner = (bp.groupby(["party", "ownership_type", bp.owner.fillna("(no owner)"), "unit"])
                .agg(stays=("unit", "size"), nights=("nights", "sum"), supplies_cost=("supplies_cost", "sum")))
    by_owner = by_owner.join(
        cl.groupby(["party", cl.ownership_type.fillna("-"), cl.owner.fillna("(no owner)"), "unit"])
          .agg(cleans=("Amount", "size"), cleaning_cost=("Amount", "sum"),
               cleans_hoa_paid=("hoa_paid", lambda s: int((s > 0).sum())), hoa_paid=("hoa_paid", "sum"),
               owner_paid=("owner_paid", "sum"))
          .rename_axis(["party", "ownership_type", "owner", "unit"]), how="outer").fillna(0).reset_index()
    by_owner["owner_total_cost"] = by_owner.supplies_cost + by_owner.owner_paid

    frac_months = (cl[cl.week_start.notna()].assign(week=lambda d: d.week_start.astype(str))
                   .groupby(["unit", "month"])
                   .agg(fraction_weeks=("week", "nunique"), cleans=("Amount", "size"),
                        not_invoiced=("Invoice #", lambda s: int((s == "not invoiced").sum())),
                        cleaning_cost=("Amount", "sum"), hoa_cleans=("hoa_clean_seq", lambda s: int((s == 1).sum())),
                        hoa_paid=("hoa_paid", "sum"), owner_paid=("owner_paid", "sum"))
                   .reset_index())
    hoa_months = (cl[cl.ownership_type == "Whole owner"].groupby(["unit", "month"])
                  .agg(cleans=("Amount", "size"), cleaning_cost=("Amount", "sum"),
                       hoa_cleans=("hoa_paid", lambda s: int((s > 0).sum())), hoa_paid=("hoa_paid", "sum"),
                       owner_cleans=("owner_paid", lambda s: int((s > 0).sum())), owner_paid=("owner_paid", "sum"))
                  .reset_index())

    # ---- exceptions
    exc = []
    for _, r in cl[cl.match == "UNMATCHED"].iterrows():
        exc.append({"type": f"clean with no checkout on or up to {max_days} days before", "unit": r.unit, "date": r.Date, "amount": r.Amount,
                    "detail": f"inv {r['Invoice #']}; paid by {r.cleaning_paid_by} (owner {r.owner})" + (f"; {r.match_note}" if isinstance(r.match_note, str) else "")})
    cleaned = set(cl.booking_idx.dropna().astype(int))
    invoiced = cl[cl["Invoice #"] != "not invoiced"].Date
    inv_first, inv_last = invoiced.min(), invoiced.max()
    for i, r in bp.iterrows():
        if i not in cleaned and inv_first <= r.check_out <= inv_last:
            exc.append({"type": "stay with no clean", "unit": r.unit, "date": r.check_out, "amount": None,
                        "detail": f"{r.guest} {r.check_in:%m/%d}-{r.check_out:%m/%d} ({r.channel}, {r.source})"})
    act = b[b.active].sort_values(["unit", "check_in"])
    for u, g in act.groupby("unit"):
        prev_out = None
        for r in g.itertuples():
            if prev_out is not None and r.check_in < prev_out[0]:
                exc.append({"type": "overlapping active stays", "unit": u, "date": r.check_in, "amount": None,
                            "detail": f"{prev_out[1]} overlaps {r.guest} {r.check_in:%m/%d}-{r.check_out:%m/%d}"})
            if prev_out is None or r.check_out > prev_out[0]:
                prev_out = (r.check_out, f"{r.guest} {r.check_in:%m/%d}-{r.check_out:%m/%d}")
    for r in bp[bp.spans_weeks == True].itertuples():  # noqa: E712
        if r.other_week_owners:
            exc.append({"type": "stay spans two owners' weeks (labelled by check-in week)", "unit": r.unit,
                        "date": r.check_in, "amount": r.supplies_cost,
                        "detail": f"{r.guest}: {r.owner} -> also {r.other_week_owners}"})
    for r in bp[bp.party == "UNASSIGNED"].groupby("unit").size().items():
        exc.append({"type": "unit not in fractional master (party UNASSIGNED)", "unit": r[0], "date": None,
                    "amount": None, "detail": f"{r[1]} stays in period"})
    exceptions = pd.DataFrame(exc)

    bookings_cols = ["unit", "check_in", "check_out", "nights", "status", "active", "in_period", "guest", "channel",
                     "prev_channel", "source", "guesty_code", "prev_res_no", "prev_alt_confirm", "week_no", "segment",
                     "week_start", "week_end", "owner", "ownership_type", "in_rental_program", "party", "is_timeshare",
                     "spans_weeks", "other_week_owners", "guesty_guests", "max_sleeps", "supply_guests",
                     "supply_guests_basis", "supplies_cost"]
    out = ROOT / "output" / f"yacinde_expense_allocation_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    combined = bookings_with_cleans(b[b.active & b.in_period], cl, master, start, end, cfg["parties"])
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        combined.to_excel(xw, sheet_name="Bookings + Cleaning", index=False)
        by_owner.to_excel(xw, sheet_name="By Owner", index=False)
        hoa_months.to_excel(xw, sheet_name="Whole Owner Cleans by Month", index=False)
        frac_months.to_excel(xw, sheet_name="Fraction Cleans by Month", index=False)
        b[bookings_cols].to_excel(xw, sheet_name="Bookings", index=False)
        lead = ["Date", "unit", "Rate", "Qty", "Amount", "cleaning_paid_by", "hoa_paid", "owner_paid", "hoa_rule", "hoa_clean_seq",
                "party", "owner", "ownership_type", "segment", "week_start", "week_end", "month"]
        cl[lead + [c for c in cl.columns if c not in lead + ["booking_idx", "clean_id", "Unit"]]].to_excel(
            xw, sheet_name="Cleaning Allocation", index=False)
        exceptions.to_excel(xw, sheet_name="Exceptions", index=False)
        dedup_log.to_excel(xw, sheet_name="Dedup Log", index=False)
        holds.to_excel(xw, sheet_name="Prev Holds Excluded", index=False)
        ws = xw.book["Bookings + Cleaning"]              # colour-code the two halves of the header
        for cell in ws[1]:
            colour = "DDEBF7" if str(cell.value).startswith("stay_") else "FCE4D6" if str(cell.value).startswith("clean_") else "E2EFDA"
            cell.fill = PatternFill("solid", fgColor=colour)
            cell.font = Font(bold=True)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "D2" if ws.title == "Bookings + Cleaning" else "A2"
            for row in ws.iter_rows(min_row=2):          # dates only, no 00:00:00 time part
                for cell in row:
                    if isinstance(cell.value, datetime.datetime):
                        cell.value = cell.value.date()
                        cell.number_format = "yyyy-mm-dd"
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(45, max(10, *(len(str(c.value or "")) for c in col[:200])) + 2)

    pd.set_option("display.width", 200)
    print(f"Guesty {len(gy)} | previous company {len(prev)} (+{len(holds)} holds dropped) -> combined {len(b)} "
          f"({b.active.sum()} active); in period {start:%Y-%m-%d}..{end:%Y-%m-%d}: {len(bp)} active stays")
    print(f"Owner weeks added from fractional master: {len(owner_weeks)} "
          + ", ".join(f"{r.unit} {r.check_in:%m/%d}-{r.check_out:%m/%d}" for r in owner_weeks.itertuples()))
    n_synth = int((cl["Invoice #"] == "not invoiced").sum())
    print(f"Cleans {len(cl) - n_synth} invoiced (${cl.Amount.sum():,.2f}) + {n_synth} $0 fraction-week HOA cleans; "
          f"unmatched {int((cl.match == 'UNMATCHED').sum())}")
    print(summary.to_string(index=False))
    print(exceptions.type.value_counts().to_string() if len(exceptions) else "no exceptions")
    print(f"-> {out}")


if __name__ == "__main__":
    main()
