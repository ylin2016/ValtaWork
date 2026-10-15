"""Payout files -> planned QBO Payments + payout JEs, for review. WRITES NOTHING TO QBO.

    python -m src.build_payouts

Two-step clearing (QBO_operations' model, owner decisions 2026-10-02 / 10-06):

    Payment     Dr Payments Clearing - <channel>  Cr A/R     one per reservation per payout
    Payout JE   Dr Chase Trust Checking 9967      Cr clearing  one per payout (bank feed matches it)

Reads the verified payout files (src.payout_inbox output):
    inputs/airbnb/airbnb_<account>_<from>_<to>.csv     one Payout row + its detail rows
    inputs/bookingcom/Payout_Statement__*.csv           grouped by Statement Descriptor
    inputs/stripe/<amount> Payout.csv                   one payout each; po_ id from the
                                                        `Stripe platform payment*.csv` list
    inputs/marriott/Marriott_Remittance_<date>.xlsx|csv one remittance advice each (owner,
                                                        2026-10-14: two-step via Payments
                                                        Clearing - Marriott, like the others)
Hopper stays manual (owner, 2026-10-06) and is not built here.

Line routing (VRP's / QBO_operations' format):
    reservation money        Cr clearing, Payment applies it to the invoice
    refund / adjustment      booking CANCELLED -> clearing; stay happened -> the part that only takes
                             back money paid beyond the invoice (Guesty's total) -> clearing (a
                             clawback, no owner effect), the rest -> Dr Resolutions (a real refund)
    money beyond invoice     -> Resolutions up to earlier refunds charged there (a re-charge);
                             the rest is held in clearing for the channel's clawback (warned)
    resolution payout        Cr Resolutions (Airbnb pays the owner side directly)
    Stripe fees              Dr Fee - Stripe Processing, split Stripe / Guesty application fee
    application fee refund   Cr Fee - Stripe Processing
    Marriott clawback        negative remittance line; same rule as a refund

Writes review/payouts_plan_<stamp>.json (what the poster will post), payouts_<stamp>.csv
(one row per payout) and payout_lines_<stamp>.csv (every Payment and JE line).
A payout with any BLOCK is planned but will not post until the cause is fixed.
"""
from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal

import yaml

from . import bridge
from .codes import invoice_doc
from .guesty_store import latest_pull
from .paths import CONFIG_DIR, INPUTS_DIR, REVIEW_DIR
from .payout_inbox import airbnb_payouts
from .resolver_util import esc
from .stripe_fees import charge_code

CFG = yaml.safe_load((CONFIG_DIR / "vrp_format.yml").read_text())
PCFG = CFG["payouts"]
START = str(yaml.safe_load((CONFIG_DIR / "payout_accounts.yml").read_text())["start_date"])
CENT = 0.005


def r2(x) -> float:
    return round(float(x) + 0.0, 2)


def money(s) -> float:
    s = str(s or "").replace(",", "").replace("$", "").strip()
    return float(Decimal(s)) if s else 0.0


def mdy(s: str) -> str:
    m, d, y = s.split("/")
    return f"{y}-{int(m):02d}-{int(d):02d}"


def read_csv(p):
    with open(p, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


# --- payout files -------------------------------------------------------------------------
def airbnb_files() -> list[dict]:
    out = []
    for p in sorted((INPUTS_DIR / "airbnb").glob("airbnb_*_*.csv")):
        acct = re.match(r"airbnb_(\d+)_", p.name).group(1)
        payouts, _ = airbnb_payouts(read_csv(p))
        for po in payouts:
            head = po["rows"][0]
            out.append({"channel": "airbnb", "id": po["ref"], "date": po["date"].isoformat(),
                        "amount": r2(po["amount"]), "account": acct, "file": p.name,
                        "bank_desc": f"Payout | {head['Details']} | {po['ref']}",
                        "rows": [{"type": r["Type"], "code": r["Confirmation code"],
                                  "amount": money(r["Amount"]), "details": r.get("Details", ""),
                                  "guest": r.get("Guest", "")} for r in po["rows"][1:]]})
    return out


def booking_files() -> list[dict]:
    by: dict[str, dict] = {}
    for p in sorted((INPUTS_DIR / "bookingcom").glob("*.csv")):
        for r in read_csv(p):
            sd = r["Statement Descriptor"]
            po = by.setdefault(sd, {"channel": "bookingcom", "id": sd, "date": r["Payout date"],
                                    "file": p.name, "rows": [],
                                    "bank_desc": f"Payout | Bank Account | {sd}"})
            po["rows"].append({"type": r["Transaction type"], "code": r["Reservation number"],
                               "amount": money(r["Total amount (gross)"]), "details": ""})
    for po in by.values():
        po["amount"] = r2(sum(x["amount"] for x in po["rows"]))
    return list(by.values())


def stripe_files() -> list[dict]:
    lists = [r for p in (INPUTS_DIR / "stripe").glob("Stripe platform payment*.csv") for r in read_csv(p)]
    out = []
    for p in sorted((INPUTS_DIR / "stripe").glob("*Payout.csv")):
        rows = read_csv(p)
        net = r2(sum(money(r["Net"]) for r in rows))
        po = [x for x in lists if abs(money(x["Amount"]) - net) < CENT]
        pid = po[0]["id"] if len(po) == 1 else None
        date = (po[0]["Arrival Date (UTC)"] or po[0]["Created (UTC)"])[:10] if pid else None
        out.append({"channel": "stripe", "id": pid, "date": date, "amount": net, "file": p.name,
                    "bank_desc": f"Payout | •••••9967 (USD) | {pid}",
                    "rows": [{"type": r["Type"], "id": r["ID"], "amount": money(r["Amount"]),
                              "fees": money(r["Fees"]), "net": money(r["Net"]),
                              "details": r.get("Description", ""),
                              "code": charge_code(r) if r["Type"] in ("Charge", "Refund") else ""}
                             for r in rows]})
    return out


def _sheet_rows(p) -> list[list[str]]:
    if p.suffix == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(p, data_only=True).active
        return [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    with p.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.reader(fh))


def _signed(s) -> float:
    s = str(s or "").strip()
    neg = s.startswith("(") and s.endswith(")")
    v = money(s.strip("()")) if s not in ("", "-") else 0.0
    return -v if neg else v


def _next_business_day(d: str) -> str:
    from datetime import date, timedelta
    x = date.fromisoformat(d[:10]) + timedelta(days=1)
    while x.weekday() >= 5:
        x += timedelta(days=1)
    return x.isoformat()


def marriott_files() -> list[dict]:
    """Marriott Payment Remittance Advice: key/value header, then one row per Marriott invoice
    (`M<13 digits><suffix>`; negative = clawback, `...PTREFUND...` = pass-through tax refund).
    The deposit lands the next business day -- staff dated their hand Payments that way."""
    out = []
    for p in sorted((INPUTS_DIR / "marriott").glob("Marriott_Remittance_*.*")):
        if p.suffix not in (".csv", ".xlsx"):
            continue
        rows, head, items, in_table = _sheet_rows(p), {}, [], False
        for r in rows:
            k = (r[0] if r else "").strip()
            if k == "Invoice ID":
                in_table = True
            elif in_table and k and k != "Total" and re.match(r"M\d{13}", k):
                amt = _signed(r[4])
                typ = "PT Refund" if "PTREFUND" in k else ("Reservation" if amt > 0 else "Clawback")
                items.append({"type": typ, "marriott_id": k, "code": k[:14], "amount": amt,
                              "details": k})
            elif not in_table and k and len(r) > 1:
                head[k] = r[1]
        ref = head.get("Payment Reference Number", "").strip()
        total = _signed(next((v for k, v in head.items() if k.startswith("Total Payment Amount")), ""))
        amount = r2(sum(x["amount"] for x in items))
        po = {"channel": "marriott", "id": f"MAR-{ref}" if ref else None, "amount": amount,
              "date": _next_business_day(head["Payment Date"]) if head.get("Payment Date") else None,
              "file": p.name, "rows": items, "bank_desc": f"Marriott remittance {ref}",
              "check": None if abs(amount - total) < CENT else f"lines sum {amount:.2f} != advice {total:.2f}"}
        out.append(po)
    return out


# --- planner ------------------------------------------------------------------------------
class Planner:
    def __init__(self, payouts: list[dict]):
        self.qbo = bridge.qbo_client()
        self.payouts = payouts
        _, pull = latest_pull()
        self.guesty = {}
        for r in pull["results"]:
            if r.get("confirmationCode"):
                self.guesty[r["confirmationCode"]] = r
                bc = ((r.get("integration") or {}).get("bookingCom") or {}).get("reservationId")
                if bc:
                    self.guesty[str(bc)] = r
        self.app_fee = self._guesty_app_fees()
        for po in payouts:                       # Marriott M<13 digits> -> Guesty HVBM-M<digits>-…
            for r in po["rows"]:
                if po["channel"] == "marriott":
                    hit = [c for c in self.guesty if c.startswith(f"HVBM-{r['code']}-")]
                    if len(hit) == 1:
                        r["code"] = hit[0]
                        continue
                    # not in the Guesty pull (unchanged since): its live invoice in QBO
                    inv = [i for i in self.qbo.query_all(
                        f"SELECT * FROM Invoice WHERE DocNumber LIKE 'HVBM-{r['code']}%'", "Invoice")
                        if "Voided" not in (i.get("PrivateNote") or "")]
                    if len(inv) == 1:
                        r["code"] = inv[0]["DocNumber"]
        # Per booking, across every payout in the run: reservation money and Airbnb
        # adjustments. When Guesty already netted an adjustment into the invoice
        # (invoice total == reservations + adjustments), the adjustment is the same money
        # coming back out of clearing -- not a refund to charge the owner again.
        self.res_sum, self.adj_sum = defaultdict(float), defaultdict(float)
        for po in payouts:
            for r in po["rows"]:
                if po["channel"] == "airbnb" and r["type"] in ("Reservation", "Pass Through Tot"):
                    self.res_sum[r["code"]] += r["amount"]
                elif po["channel"] == "airbnb" and r["type"] == "Adjustment":
                    self.adj_sum[r["code"]] += r["amount"]
        self._load_invoices()
        self._load_existing()
        # Per (channel, customer), changes planned in this run: money held in clearing beyond the
        # invoice (awaiting a clawback) and refunds charged to Resolutions. QBO's prior state is
        # added by _prior().
        self.held, self.res_refunded, self._prior_cache = defaultdict(float), defaultdict(float), {}

    def _guesty_app_fees(self) -> dict[str, float]:
        """Stripe charge id -> the Guesty application fee actually taken (from Guesty)."""
        from .stripe_fees import _charges
        out = {}
        for r in self.guesty.values():
            for p in (r.get("money") or {}).get("payments") or []:
                for c in _charges(p):
                    if c.get("application_fee_amount") is not None:
                        out[c["id"]] = c["application_fee_amount"] / 100.0
        return out

    def _doc(self, channel: str, code: str) -> str:
        if channel == "bookingcom":
            return code[:21]
        g = self.guesty.get(code)
        return invoice_doc(g) if g else code[:21]

    def _load_invoices(self):
        docs = sorted({self._doc(po["channel"], r["code"]) for po in self.payouts
                       for r in po["rows"] if r.get("code") and not r["code"].startswith("Application fee")})
        self.invoices = {}
        for i in range(0, len(docs), 40):
            part = ", ".join(f"'{esc(d)}'" for d in docs[i:i + 40])
            for inv in self.qbo.query_all(f"SELECT * FROM Invoice WHERE DocNumber IN ({part})", "Invoice"):
                self.invoices[inv["DocNumber"]] = inv

    def _load_existing(self):
        # from a week before the earliest payout (Marriott remittances predate START)
        from datetime import date, timedelta
        d0 = min([START] + [po["date"] for po in self.payouts if po.get("date")])
        since = (date.fromisoformat(d0) - timedelta(days=7)).isoformat()
        je_post = bridge._sub("qbo_ops", bridge.QBO_OPS_ROOT, "je.post")
        self.bank_index = je_post.existing_payouts(self.qbo, start=since)
        self.je_docs = {j.get("DocNumber") for j in self.qbo.query_all(
            f"SELECT * FROM JournalEntry WHERE TxnDate >= '{since}'", "JournalEntry")}
        self.payments = self.qbo.query_all(f"SELECT * FROM Payment WHERE TxnDate >= '{since}'", "Payment")

    # -- per invoice facts
    def facts(self, channel, code):
        inv = self.invoices.get(self._doc(channel, code))
        g = self.guesty.get(code)
        if not inv:
            return None, g
        ln = next((l for l in inv["Line"] if l.get("DetailType") == "SalesItemLineDetail"), None)
        span = (ln.get("Description") or "").split(" | ", 1)[1] if ln and " | " in (ln.get("Description") or "") else ""
        klass = ((ln or {}).get("SalesItemLineDetail", {}).get("ClassRef") or {}).get("name")
        if not klass and g:              # a voided invoice keeps no class: class + dates from Guesty
            klass, span = self._guesty_class(g), self._guesty_span(g)
        return {"Id": inv["Id"], "doc": inv["DocNumber"], "balance": r2(inv["Balance"]),
                "total": r2(inv["TotalAmt"]), "customer": inv["CustomerRef"]["name"],
                "customer_id": inv["CustomerRef"]["value"], "class": klass,
                "span": span, "voided": "Voided" in (inv.get("PrivateNote") or "")}, g

    def _guesty_class(self, g) -> str | None:
        """Live class FQN as build_invoices resolves it: statement class map, else Listings:<nick>."""
        if not hasattr(self, "_res"):
            self._res, self._cmap = bridge.resolver(self.qbo), bridge.class_map()
        nick = (g.get("listing") or {}).get("nickname") or ""
        pid = bridge.to_property_id(nick)
        for c in filter(None, [(self._cmap.get(pid) or {}).get("qbo_class_name"), f"Listings:{nick}"]):
            fqn = self._res.klass_fqn(c)
            if fqn:
                return fqn
        return None

    @staticmethod
    def _guesty_span(g) -> str:
        ci, co = (g.get("checkIn") or "")[:10], (g.get("checkOut") or "")[:10]
        return f"{ci} to {co} | {g.get('nightsCount', '')} nights" if ci else ""

    def _prior(self, ch: str, f: dict) -> tuple[float, float]:
        """What QBO already holds for this reservation's customer, before this run:
        (money in the channel clearing beyond what Payments applied -- an overpayment awaiting
        its clawback, refunds charged to Resolutions not yet credited back)."""
        key = (ch, f["customer_id"])
        if key not in self._prior_cache:
            if not hasattr(self, "_jes"):
                self._jes = self.qbo.query_all("SELECT * FROM JournalEntry WHERE TxnDate >= '2026-01-01'",
                                               "JournalEntry")
                self._acct_ids = {}
            clearing, res = PCFG["clearing"][ch], PCFG["resolutions"]
            cr = rs = 0.0
            for j in self._jes:
                for l in j["Line"]:
                    d = l.get("JournalEntryLineDetail") or {}
                    if ((d.get("Entity") or {}).get("EntityRef") or {}).get("value") != f["customer_id"]:
                        continue
                    sign = 1 if d["PostingType"] == "Credit" else -1
                    name = d["AccountRef"].get("name") or ""
                    if name.endswith(clearing.split(":")[-1]):
                        cr += sign * l["Amount"]
                    elif name.endswith(res.split(":")[-1]):
                        rs -= sign * l["Amount"]             # net debit = refunds charged
            if clearing not in self._acct_ids:
                a = self.qbo.query_all(f"SELECT * FROM Account WHERE Name = '{esc(clearing.split(':')[-1])}'", "Account")
                self._acct_ids[clearing] = a[0]["Id"] if a else None
            paid = sum(float(x["TotalAmt"]) for x in self.qbo.query_all(
                f"SELECT * FROM Payment WHERE CustomerRef = '{f['customer_id']}'", "Payment")
                if (x.get("DepositToAccountRef") or {}).get("value") == self._acct_ids[clearing])
            self._prior_cache[key] = (max(r2(cr - paid), 0.0), max(r2(rs), 0.0))
        return self._prior_cache[key]

    def split_refund(self, ch, code, f, amt, pay_by_code, applied) -> tuple[float, float]:
        """A negative row on a stay that happened -> (to clearing, to Resolutions), both positive.
        Only money paid beyond the invoice can be clawed back without touching the owner: first
        this payout's own excess over the open balance, then what earlier payouts left in clearing.
        Booking.com labels a clawback and a guest refund the same (REFUND); this tells them apart
        (Castano 6314669620: paid 3,395.77 on a 1,805.82 stay, 1,589.95 clawed back next day)."""
        key, back = (ch, f["customer_id"]), r2(-amt)
        room = max(r2(f["balance"] - applied[f["Id"]]), 0.0)
        in_po = min(back, max(r2(pay_by_code[code] - room), 0.0))
        pay_by_code[code] -= in_po                    # the Payment stays at the open balance
        held = min(r2(back - in_po), max(r2(self._prior(ch, f)[0] + self.held[key]), 0.0))
        self.held[key] -= held
        to_res = r2(back - in_po - held)
        self.res_refunded[key] += to_res
        return r2(in_po + held), to_res

    def netted(self, code: str, f: dict) -> bool:
        a = self.adj_sum.get(code, 0.0)
        return a < -CENT and abs(self.res_sum.get(code, 0.0) + a - f["total"]) < 0.02

    def plan(self) -> list[dict]:
        plans, applied = [], defaultdict(float)
        for po in sorted(self.payouts, key=lambda x: (x["channel"], x["date"] or "", x["id"] or "")):
            plans.append(self._plan_one(po, applied))
        # A cancelled Marriott booking's money and clawbacks wash in clearing; anything left
        # over (Marriott took back more than it paid, or a remittance is missing) is flagged.
        net, last = defaultdict(float), {}
        for po, p in zip(sorted(self.payouts, key=lambda x: (x["channel"], x["date"] or "", x["id"] or "")), plans):
            if po["channel"] == "marriott" and p["action"] != "SKIP":
                for r in po["rows"]:
                    f, g = self.facts("marriott", r["code"])
                    if f and ((g or {}).get("status") == "canceled" or f["voided"]):
                        net[r["code"]] += r["amount"]
                        last[r["code"]] = p
        for code, n in net.items():
            if abs(n) > CENT:
                last[code]["warn"].append(f"{code} is CANCELLED and its remittances net {n:+.2f} in "
                                          f"Payments Clearing - Marriott -- not zero; ask Marriott / decide")
        return plans

    def _plan_one(self, po, applied) -> dict:
        ch = po["channel"]
        p = {**{k: po[k] for k in ("channel", "id", "date", "amount", "file")},
             "payments": [], "lines": [], "block": [], "warn": [], "action": "CREATE"}
        if not po["id"] or not po["date"]:
            p["block"].append("no payout id/date (Stripe file not in the payout list?)")
        if po.get("check"):
            p["block"].append(f"remittance does not add up: {po['check']}")
        if po["date"] and po["date"] < START:
            # VRP era -- and for Marriott (never VRP's) the hand-entered era: older remittances
            # sit in inputs/marriott/ for review only and are already in the books (cutover 9/28)
            p["action"] = "SKIP"
            p["warn"].append(f"before {START} ({'hand-entered' if ch == 'marriott' else 'VRP'} era)")
            return p
        doc = (po["id"] or "")[:21]
        p["DocNumber"], p["memo"] = doc, f"External Payment ID: {po['id']}"
        if doc in self.je_docs or (po["date"], po["amount"]) in self.bank_index:
            p["action"] = "SKIP"
            hit = self.bank_index.get((po["date"], po["amount"]))
            p["warn"].append(f"already booked in QBO ({hit[0] if hit else doc})")
            return p
        hand = self._hand_recorded(po)
        if hand:
            p["action"] = "SKIP"
            p["warn"].append(f"already recorded by hand from the bank feed ({hand}) -- not booked again")
            return p
        clearing = PCFG["clearing"][ch]
        pay_by_code = defaultdict(float)

        def line(acct, amt, desc, f=None, cls=None, location=None):
            # amt > 0 = credit, < 0 = debit (the payout's sign convention)
            p["lines"].append({"account": acct, "posting": "Credit" if amt >= 0 else "Debit",
                               "amount": r2(abs(amt)), "description": desc,
                               "customer": (f or {}).get("customer"),
                               "customer_id": (f or {}).get("customer_id"),
                               "class": cls if cls is not None else (f or {}).get("class"),
                               "location": location})

        def refund(amt, desc, f, code, cancelled):
            # cancelled -> washes in clearing; else clawback of an overpayment -> clearing,
            # the rest (a real refund on a stay that happened) -> Resolutions
            if cancelled:
                return line(clearing, amt, desc, f)
            claw, res = self.split_refund(ch, code, f, amt, pay_by_code, applied)
            if claw:
                line(clearing, -claw, desc + (" | clawback of an overpayment" if res else ""), f)
                p["warn"].append(f"{code}: {claw:.2f} of the {-amt:.2f} refund takes back money paid "
                                 f"beyond invoice {f['doc']} -> clearing, not Resolutions")
            if res:
                line(PCFG["resolutions"], -res, desc, f)

        for r in po["rows"]:
            typ, code, amt = r["type"], r.get("code") or "", r["amount"]
            if ch == "stripe" and typ == "Adjustment" and r["details"].startswith("Chargeback"):
                # Chargeback withdrawal / reversal: the disputed money leaves or comes back
                # through clearing, on the reservation it was charged for -- as the 2026-08
                # withdrawals were booked (JE 116571). Old charges are found by the line that
                # booked their withdrawal.
                m = re.search(r"(ch_\w+)", r["details"])
                f = self._facts_for_charge(m.group(1)) if m else None
                if not f and m:
                    f = self._facts_from_withdrawal(m.group(1))
                if not f:
                    p["block"].append(f"{r['details']} {amt:.2f}: cannot find its reservation -- ask")
                    continue
                line(clearing, amt, f"{r['details']} | {self._tail(ch, f)}", f)
                continue
            if ch == "stripe" and typ == "Stripe Fee":
                # Account-level fees (dispute counter fees, refunded when the dispute is won):
                # Legal and Professional Fees, company class/location, no customer (JE 116571).
                line(PCFG["dispute_fee"], amt, r["details"], None, PCFG["company_class"],
                     PCFG["company_location"])
                continue
            if ch == "stripe" and typ == "Adjustment":
                m = re.search(r"(ch_\w+)", r["details"])
                src = next((x for x in po["rows"] if m and x.get("id") == m.group(1)), None)
                f, _ = self.facts(ch, src["code"]) if src else (None, None)
                if not f:                       # the charge itself sits in an earlier payout
                    f = self._facts_for_charge(m.group(1) if m else "")
                line(PCFG["stripe_fee"], amt, f"{r['details']} | {self._tail(ch, f)}", f)
                continue
            f, g = self.facts(ch, code)
            if not f:
                p["block"].append(f"{typ} {code} {amt:.2f}: no invoice in QBO")
                continue
            cancelled = (g or {}).get("status") == "canceled" or f["voided"]
            tail = self._tail(ch, f)
            if ch == "stripe" and typ == "Charge":
                line(clearing, amt, f"{code} | {tail}", f)
                guesty_fee = r2(self.app_fee.get(r["id"], round(amt * 0.01, 2)))
                stripe_fee = r2(r["fees"] - guesty_fee)
                line(PCFG["stripe_fee"], -stripe_fee, f"Stripe processing fees | {tail}", f)
                line(PCFG["stripe_fee"], -guesty_fee, f"Guesty application fee | {tail}", f)
                if not cancelled:
                    pay_by_code[code] += amt
            elif ch == "stripe" and typ == "Refund":
                refund(amt, f"REFUND FOR CHARGE ({code}) | {tail}", f, code, cancelled)
            elif ch == "airbnb" and typ in ("Reservation", "Pass Through Tot"):
                line(clearing, amt, f"{typ} | airbnb | {f['span']}" + (f" | {r['details']}" if r["details"] else ""), f)
                if not cancelled:
                    pay_by_code[code] += amt
            elif ch == "airbnb" and typ == "Pass Through Tot Adjustment" and not cancelled:
                p["block"].append(f"no rule for airbnb {typ!r} on a stay that happened ({code} {amt:.2f}) -- ask")
                continue
            elif ch == "airbnb" and typ in ("Adjustment", "Resolution Payout", "Resolution Adjustment",
                                            "Pass Through Tot Adjustment"):
                # A cancelled booking's clawbacks (incl. a mid-stay cancellation's resolution
                # adjustment) wash in clearing: the invoice drops to the kept amount at month end,
                # so charging Resolutions too would hit the owner twice (owner, 2026-10-14).
                acct = (clearing if (typ == "Adjustment" and self.netted(code, f))
                        or (cancelled and typ != "Resolution Payout") else PCFG["resolutions"])
                line(acct, amt, f"{typ} | airbnb | {f['span']}" + (f" | {r['details']}" if r["details"] else ""), f)
            elif ch == "airbnb" and typ == "Misc Credit":
                line(PCFG["commission_revenue"], amt, f"{typ} | airbnb | {r['details']}", f, PCFG["company_class"])
            elif ch == "marriott" and typ == "Reservation":
                line(clearing, amt, f"{r['details']} | marriott | {f['span']}", f)
                if not cancelled:
                    pay_by_code[code] += amt
            elif ch == "marriott" and typ in ("Clawback", "PT Refund"):
                refund(amt, f"{typ} {r['details']} | marriott | {f['span']}", f, code, cancelled)
            elif ch == "bookingcom" and typ == "RESERVATION":
                line(clearing, amt, f"RESERVATION | bookingcom | {f['span']}", f)
                if not cancelled:
                    pay_by_code[code] += amt
            elif ch == "bookingcom" and typ == "REFUND":
                refund(amt, f"REFUND | bookingcom | {f['span']}", f, code, cancelled)
            else:
                p["block"].append(f"no rule for {ch} row type {typ!r} ({code} {amt:.2f}) -- ask")
                continue
            if cancelled and typ in ("Charge", "Reservation", "RESERVATION") and amt > 0:
                p["warn"].append(f"{code} is CANCELLED: {amt:.2f} credited to clearing, no Payment "
                                 f"(its refund washes there); existing payments on the invoice need voiding")

        # the bank line: the whole payout
        line(PCFG["bank"], -po["amount"], po["bank_desc"])
        dr = sum(l["amount"] for l in p["lines"] if l["posting"] == "Debit")
        cr = sum(l["amount"] for l in p["lines"] if l["posting"] == "Credit")
        if abs(dr - cr) > CENT:
            p["block"].append(f"JE does not balance: debits {dr:.2f} credits {cr:.2f}")

        # payments, one per reservation in this payout
        for code, amt in pay_by_code.items():
            f, _ = self.facts(ch, code)
            dup = [x for x in self.payments if x["CustomerRef"]["value"] == f["customer_id"]
                   and po["id"] and po["id"] in (x.get("PrivateNote") or "")]
            if dup:
                p["warn"].append(f"{code}: payment for {po['id']} already exists (Id {dup[0]['Id']}) -- skipped")
                continue
            room = r2(f["balance"] - applied[f["Id"]])
            if amt > room + CENT:
                excess = r2(amt - max(room, 0.0))
                amt = r2(max(room, 0.0))
                if ch == "airbnb" and self.netted(code, f):
                    p["warn"].append(f"{code}: payment capped at the open balance {room:.2f}; the "
                                     f"{excess:.2f} over it comes back out of clearing with its Airbnb "
                                     f"adjustment (already netted on the invoice)")
                else:
                    # Money beyond what the invoice asks for: a re-charge after a refund charged to
                    # Resolutions goes back there; anything more is the channel overpaying -- held
                    # in clearing for its clawback (Castano 6314669620), never the owner's income.
                    key = (ch, f["customer_id"])
                    back = min(excess, max(r2(self._prior(ch, f)[1] + self.res_refunded[key]), 0.0))
                    if back:
                        self.res_refunded[key] -= back
                        line(clearing, -back, f"Excess over invoice {f['doc']} -> Resolutions | {self._tail(ch, f)}", f)
                        line(PCFG["resolutions"], back, f"Received beyond invoice {f['doc']} | {self._tail(ch, f)}", f)
                        p["warn"].append(f"{code}: {back:.2f} received beyond invoice {f['doc']} (open "
                                         f"{room:.2f}) after a refund charged there -> credited to Resolutions")
                    hold = r2(excess - back)
                    if hold:
                        self.held[key] += hold
                        p["warn"].append(f"{code}: {hold:.2f} paid beyond invoice {f['doc']} (open {room:.2f}) "
                                         f"held in clearing for the channel's clawback -- if none comes by "
                                         f"month end, check the invoice against Guesty")
            if amt < CENT:
                continue
            applied[f["Id"]] += amt
            p["payments"].append({"code": code, "invoice_id": f["Id"], "invoice_doc": f["doc"],
                                  "customer": f["customer"], "customer_id": f["customer_id"],
                                  "amount": r2(amt), "TxnDate": po["date"], "PaymentRefNum": code[:21],
                                  "PrivateNote": p["memo"], "deposit_to": clearing})
        if p["block"]:
            p["action"] = "BLOCKED"
        return p

    def _tail(self, ch, f) -> str:
        src = {"stripe": "", "airbnb": "airbnb", "bookingcom": "bookingcom", "marriott": "marriott"}[ch]
        if ch == "stripe":
            src = (f["customer"] or "").split(" - ")[0]
        return f"{src} | {f['span']}"

    def _hand_recorded(self, po) -> str | None:
        """Staff record some payouts straight from the bank feed: a Payment of the whole payout
        on the reservation's invoice, without our `External Payment ID` memo. When every
        reservation row of the payout has such a Payment (same amount, within 5 days), the
        payout is already in the books -- a JE would count the deposit twice."""
        money = ("Reservation", "Pass Through Tot", "RESERVATION", "Charge")
        res = [r for r in po["rows"] if r["type"] in money]
        if not res or len(res) != len(po["rows"]) and po["channel"] != "stripe":
            return None
        by_code = defaultdict(float)                 # a reservation's money across its rows
        for r in res:
            by_code[r["code"]] += r["amount"]
        d0 = datetime.fromisoformat(po["date"]).date()
        ids = []
        for code, amt in by_code.items():
            f, _ = self.facts(po["channel"], code)
            hit = f and next((x for x in self.payments
                              if x["CustomerRef"]["value"] == f["customer_id"]
                              and "External Payment ID" not in (x.get("PrivateNote") or "")
                              and abs(float(x["TotalAmt"]) - amt) < CENT
                              and abs((datetime.fromisoformat(x["TxnDate"]).date() - d0).days) <= 5), None)
            if not hit:
                return None
            ids.append(f"Payment {hit['Id']}")
        return ", ".join(ids)

    def _facts_from_withdrawal(self, chid: str):
        """The reservation whose chargeback withdrawal of `chid` is already booked in QBO."""
        if not hasattr(self, "_po_jes"):
            self._po_jes = self.qbo.query_all("SELECT * FROM JournalEntry WHERE TxnDate >= '2025-01-01' "
                                              "AND DocNumber LIKE 'po_%'", "JournalEntry")
        for j in self._po_jes:
            for l in j["Line"]:
                if f"withdrawal for {chid}" in (l.get("Description") or ""):
                    ent = (l["JournalEntryLineDetail"].get("Entity") or {}).get("EntityRef") or {}
                    code = ent.get("name", "").rsplit(" - ", 1)[-1]
                    if code and code not in self.invoices:      # not in this run's payout files
                        for inv in self.qbo.query_all(
                                f"SELECT * FROM Invoice WHERE DocNumber = '{esc(code[:21])}'", "Invoice"):
                            self.invoices[inv["DocNumber"]] = inv
                    f, _ = self.facts("stripe", code)
                    if f:
                        return f
        return None

    def _facts_for_charge(self, chid: str):
        for r in self.guesty.values():
            from .stripe_fees import _charges
            for pmt in (r.get("money") or {}).get("payments") or []:
                if any(c["id"] == chid for c in _charges(pmt)):
                    return self.facts("stripe", r["confirmationCode"])[0]
        return None


# --- output -------------------------------------------------------------------------------
def write(plans: list[dict], stamp: str) -> None:
    REVIEW_DIR.mkdir(exist_ok=True)
    (REVIEW_DIR / f"payouts_plan_{stamp}.json").write_text(json.dumps(plans, indent=1, default=str))
    with (REVIEW_DIR / f"payouts_{stamp}.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["action", "channel", "payout_id", "date", "amount", "payments", "payments_total",
                    "je_lines", "file", "block", "warn"])
        for p in plans:
            w.writerow([p["action"], p["channel"], p["id"], p["date"], p["amount"], len(p["payments"]),
                        r2(sum(x["amount"] for x in p["payments"])), len(p["lines"]), p["file"],
                        " | ".join(p["block"]), " | ".join(p["warn"])])
    with (REVIEW_DIR / f"payout_lines_{stamp}.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["payout_id", "date", "object", "account_or_invoice", "debit", "credit", "customer",
                    "class", "description"])
        for p in plans:
            if p["action"] == "SKIP":
                continue
            for x in p["payments"]:
                w.writerow([p["id"], x["TxnDate"], "Payment", f"Invoice {x['invoice_doc']} -> {x['deposit_to'].split(':')[-1]}",
                            x["amount"], "", x["customer"], "", x["PrivateNote"]])
            for l in p["lines"]:
                w.writerow([p["id"], p["date"], "JE", l["account"].split(":")[-1],
                            l["amount"] if l["posting"] == "Debit" else "",
                            l["amount"] if l["posting"] == "Credit" else "",
                            l["customer"] or "", l["class"] or "", l["description"]])


def main():
    payouts = airbnb_files() + booking_files() + stripe_files() + marriott_files()
    plans = Planner(payouts).plan()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    write(plans, stamp)
    from collections import Counter
    print(f"{len(plans)} payout(s):", dict(Counter((p['channel'], p['action']) for p in plans)))
    for p in plans:
        if p["action"] != "SKIP":
            print(f"  {p['action']:8} {p['channel']:10} {p['date']} {str(p['id']):22} {p['amount']:>10,.2f}"
                  f"  payments {len(p['payments'])} ({sum(x['amount'] for x in p['payments']):,.2f})")
            for b in p["block"]:
                print(f"      BLOCK {b}")
            for wn in p["warn"]:
                print(f"      warn  {wn}")
    print(f"-> review/payouts_{stamp}.csv, payout_lines_{stamp}.csv, payouts_plan_{stamp}.json")


if __name__ == "__main__":
    main()
