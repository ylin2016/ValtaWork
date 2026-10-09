"""Per-listing action report.
  output/listing_actions_2026-10-06.pdf   summary + one block per listing
  output/listing_actions_2026-10-06.xlsx  action tracker (one row per action)
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"
A = pd.read_pickle(ROOT / "data" / "actions.pkl")
PCOL = {"P1": "#E53935", "P2": "#FB8C00", "P3": "#FDD835", "OK": "#43A047"}
PTXT = {"P1": "P1 · this week", "P2": "P2 · this month", "P3": "P3 · confirm", "OK": "OK · no change"}
CHANGES = [
    "<b>Market check now uses each listing's own Wheelhouse comp set</b>, at every lead time: its booked ADR against "
    "last year where enough is on the books, otherwise its last-60-days trend against last year. It now covers 93% of "
    "nights (was 7%). Where a comp set has genuinely dropped, floors come down with it; the Eastside 2BR set is the "
    "clearest case (ADR $330 vs $402 last year).",
    "<b>New comp-set step</b> between portfolio siblings and the neighborhood band: what the comp set booked around "
    "that date last year, scaled to how this listing books against its set. Nights resting on the neighborhood band "
    "(comps' asking prices) fell from 25% to 8%.",
    "<b>Pacing added</b> from Wheelhouse neighborhood occupancy (booked so far vs typical at this lead time): well "
    "behind lowers the floor 10%, well ahead lifts the ceiling 10%.",
    "<b>Result:</b> 4,266 flagged nights across the portfolio (was 5,146). 11 listings changed; every change "
    "softened a call (Bellevue 10409, Cottage 7 and Elektra 1004 moved P1 → P2; Elektra 1413 is now OK). Longbranch, "
    "the OSBR minimums and the minimum-setting gaps are unchanged.",
]
THEME_ORDER = ["Raise the price level", "Raise the minimum", "Yacinde minimum decision", "Minimum not enforced",
               "Fix specific nights", "Not in Wheelhouse / closed", "Confirm a setting", "No change"]
THEME_WHY = {
    "Raise the price level": "Under-floor on many nights vs what it booked LY and what's already booked this season. "
                             "One Wheelhouse base-adjustment change fixes all of them.",
    "Raise the minimum": "Wheelhouse wants to go lower; price sits on the minimum most nights. Raise the minimum to "
                         "the weakest quarter of what it booked last Oct–Mar.",
    "Yacinde minimum decision": "Minimums of $300–400 but nights post at $180–250. Decide whether the minimum is an "
                                "owner/HOA requirement (enforce it) or a placeholder (lower it).",
    "Minimum not enforced": "Wheelhouse posts below the listing's own minimum, most likely a lower date-range "
                            "minimum. Either delete the override or update the default to match.",
    "Fix specific nights": "Level is fine; a handful of nights are clearly off. Fix them by hand.",
    "Not in Wheelhouse / closed": "Outside Wheelhouse or no nights for sale — nobody is checking these prices.",
    "Confirm a setting": "Usually a leftover −10% adjustment or a new listing priced well below comps.",
    "No change": "Inside the band.",
}


def money(x):
    return "—" if x is None or (isinstance(x, float) and np.isnan(x)) else f"${x:,.0f}"


def stat_line(r):
    adj = "—" if np.isnan(r.adjustment) else f"{r.adjustment:.2f}"
    return (f"LY booked {money(r.ly_adr)} · booked this season {money(r.fwd_booked_adr)} · asking now "
            f"{money(r.ask_median)} · comps {money(r.comp_median)} · WH min {money(r.min_price)} · adj {adj} · "
            f"{r.nights_avail} nights open")


def pdf(path):
    st = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=st["BodyText"], fontSize=8.5, leading=11)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#546E7A"))
    bullet = ParagraphStyle("u", parent=body, leftIndent=10, bulletIndent=2)
    h1 = ParagraphStyle("h1", parent=st["Title"], fontSize=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=st["Heading2"], fontSize=12.5, spaceBefore=8, spaceAfter=4)
    name_st = ParagraphStyle("n", parent=body, fontSize=10.5, leading=13, fontName="Helvetica-Bold")
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.55 * inch, rightMargin=0.55 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.5 * inch,
                            title="Listing pricing actions", author="Valta Realty")
    S = [Paragraph("Pricing actions by listing", h1),
         Paragraph("Run Oct 6, 2026 · next 180 nights · all 84 active Guesty listings · "
                   "sources: Guesty confirmed bookings + calendar, Wheelhouse comps/market/settings", small),
         Spacer(1, 8)]

    # priority counts
    pc = A.priority.value_counts()
    cells = [[Paragraph(f"<b>{PTXT[p]}</b><br/>{pc.get(p, 0)} listings", ParagraphStyle(
        "pc", parent=body, textColor=colors.white if p != "P3" else colors.black, alignment=1))
        for p in ("P1", "P2", "P3", "OK")]]
    t = Table(cells, colWidths=[1.82 * inch] * 4, rowHeights=[0.5 * inch])
    t.setStyle(TableStyle([("BACKGROUND", (i, 0), (i, 0), colors.HexColor(PCOL[p])) for i, p in
                           enumerate(("P1", "P2", "P3", "OK"))] + [("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    S += [t, Spacer(1, 10)]

    # what's driving it
    n = A.theme.value_counts()
    neg = int((A.adjustment == 0.9).sum())
    below = int(A.min_pattern.isin(["Systemic", "Partial"]).sum())
    S.append(Paragraph("What's driving it", h2))
    drivers = [
        f"<b>{n.get('Raise the price level', 0)} listings are priced 20–40% below what they booked last year</b> — and "
        "in most cases stays already booked this season came in even higher. Each gets a specific new Wheelhouse base "
        "adjustment (e.g. Redmond 14707 1.00 → 1.35).",
        f"<b>{n.get('Raise the minimum', 0)} listings sit on their minimum price most nights</b>, including the OSBR "
        "cottages. Each gets a new minimum set at the cheapest quarter of its stays last Oct–Mar (e.g. Cottage 5 "
        "$99 → $125).",
        f"<b>{n.get('Yacinde minimum decision', 0)} Yacinde units need one decision:</b> is the $300–400 minimum an "
        "owner/HOA requirement, or an onboarding placeholder? The answer settles all of them.",
        f"<b>{below} listings are being priced below their own minimum.</b> Each block says whether to remove the "
        "lower price setting or lower the minimum, based on what it booked last year and what's already booked "
        "this season.",
        "<b>3 listings aren't in Wheelhouse</b> (Beachwood 1, Sammamish 5124-1, the all-cottages bundle), so nobody "
        "is checking their prices.",
        f"<b>{neg} listings still carry a −10% adjustment</b> — confirm each was deliberate.",
        "Each block also lists specific nights to fix by hand. Where a nightly flag contradicts the listing-wide fix, "
        "the block says to ignore it.",
    ]
    S += [Paragraph(x, bullet, bulletText="•") for x in drivers]
    caution = Table([[Paragraph(
        "<b>Before you act:</b> the 'lower price setting' behind below-minimum prices is inferred. The Wheelhouse API "
        "only shows each listing's default minimum. Check one listing in Wheelhouse first — Mercer 3925 sits at $100 "
        "against a $129 minimum — before deleting anything across the portfolio. Prices may also have moved since the "
        "Oct 6 morning data pull. Nothing has been changed in Wheelhouse or Guesty.", body)]],
        colWidths=[7.4 * inch])
    caution.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF8E1")),
                                 ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#FB8C00")),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                                 ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    S += [Spacer(1, 6), caution, Spacer(1, 6)]
    S.append(Paragraph("What changed in this version", h2))
    S += [Paragraph(x, bullet, bulletText="•") for x in CHANGES]

    # themes
    S.append(Paragraph("What's going on — by root cause", h2))
    rows = [["Root cause", "Listings", "Why / fix"]]
    for th in THEME_ORDER:
        g = A[A.theme == th]
        if g.empty:
            continue
        names = ", ".join(g.listing.tolist())
        rows.append([Paragraph(f"<b>{th}</b>", body), Paragraph(f"{len(g)}<br/><font size=7>{names}</font>", body),
                     Paragraph(THEME_WHY[th], body)])
    t = Table(rows, colWidths=[1.35 * inch, 3.1 * inch, 2.95 * inch], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#263238")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0BEC5")),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7F8")])]))
    S += [t, Spacer(1, 8)]

    S.append(Paragraph("Do these first", h2))
    p1lvl = A[(A.priority == "P1") & (A.theme == "Raise the price level")].sort_values("severe_down", ascending=False)
    first = [
        f"<b>Audit date-range minimums in Wheelhouse.</b> {below} listings post below their own minimum, most sitting on "
        "one round number for weeks, the signature of a lower date-range minimum the API doesn't expose. Open "
        "Mercer 3925 (posts $100 vs $129 min) to find where it's set, then fix each listing per its block below.",
        "<b>Settle the Yacinde minimum question</b> (10 units, $300–400 minimums vs ~$200–220 posted): owner/HOA "
        "requirement or placeholder? The answer decides all 10 at once.",
        f"<b>Raise the level on the {len(p1lvl)} P1 'raise the price level' listings</b>. Biggest gaps: "
        + ", ".join(p1lvl.listing.head(8)) + ". Each is asking 20–40% under what it booked last year, and stays "
        "already booked this season mostly came in higher still.",
        "<b>Raise the minimums on the 'raise the minimum' listings</b>, mostly OSBR cottages: they sit on the minimum "
        "most nights.",
        f"<b>Clear the {neg} leftover −10% adjustments</b> unless each was deliberate.",
        "<b>Re-run this report after the changes</b> to see what moved and fill in the 'was' prices.",
    ]
    S += [Paragraph(f"{i}. {x}", bullet) for i, x in enumerate(first, 1)]
    S.append(Spacer(1, 6))
    S.append(Paragraph(
        "How to read a block: <b>LY booked</b> = median nightly rate of last Oct–Mar stays · <b>booked this season</b> "
        "= median of stays already on the books for the next 180 days · <b>asking now</b> = median Guesty price on "
        "open nights · <b>comps</b> = Wheelhouse neighborhood median · <b>adj</b> = Wheelhouse base adjustment. "
        "comps is the neighborhood median of comps' asking prices, so it runs above what units actually book. "
        "Floors/ceilings for single dates follow the band method in the pricing-flags report. 'Lower date-range "
        "minimum' is inferred from the posted prices — confirm it in Wheelhouse before deleting anything.", small))

    # per-listing
    for p in ("P1", "P2", "P3", "OK"):
        g = A[A.priority == p]
        if g.empty:
            continue
        S += [PageBreak(), Paragraph(f"{PTXT[p]} — {len(g)} listings", h2)]
        for r in g.itertuples():
            chip = Table([[Paragraph(f"<b>{p}</b>", ParagraphStyle(
                "c", parent=body, alignment=1, textColor=colors.white if p != "P3" else colors.black))]],
                colWidths=[0.42 * inch], rowHeights=[0.24 * inch])
            chip.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(PCOL[p])),
                                      ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            head = Table([[chip, Paragraph(f"{r.listing} <font size=8 color='#546E7A'>· {r.city or ''} · "
                                           f"{'' if pd.isna(r.bedrooms) else int(r.bedrooms)}BR · {r.theme}</font>",
                                           name_st)]], colWidths=[0.5 * inch, 6.9 * inch])
            head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
            block = [head, Paragraph(f"<b>{r.headline}</b>", body), Paragraph(stat_line(r), small)]
            block += [Paragraph(a, bullet, bulletText="•") for a in r.actions]
            block.append(Spacer(1, 9))
            S.append(KeepTogether(block))
    doc.build(S, onLaterPages=_footer, onFirstPage=_footer)


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#90A4AE"))
    canvas.drawString(0.55 * inch, 0.3 * inch, "Pricing actions by listing · Oct 6, 2026")
    canvas.drawRightString(letter[0] - 0.55 * inch, 0.3 * inch, f"Page {doc.page}")
    canvas.restoreState()


def xlsx(path):
    rows = []
    for r in A.itertuples():
        for i, a in enumerate(r.actions, 1):
            rows.append({"priority": r.priority, "listing": r.listing, "root_cause": r.theme, "headline": r.headline,
                         "#": i, "action": a, "owner": "", "status": "Open", "done_on": "", "notes": ""})
    T = pd.DataFrame(rows)
    S = A[["priority", "listing", "city", "bedrooms", "theme", "headline", "ly_adr", "fwd_booked_adr", "ask_median",
           "comp_median", "min_price", "adjustment", "min_pattern", "nights_below_min", "nights_avail",
           "down_B", "down_M", "down_I", "up_B", "up_M", "up_I"]].rename(columns={
        "theme": "root_cause", "ly_adr": "LY_booked_median", "fwd_booked_adr": "booked_this_season_median",
        "ask_median": "asking_now_median", "comp_median": "comp_median", "min_price": "wh_min_price",
        "adjustment": "wh_adjustment"})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        T.to_excel(xw, sheet_name="action_tracker", index=False)
        S.to_excel(xw, sheet_name="listing_summary", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
            hdr = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2):
                pcell = row[hdr.index("priority")]
                pcell.fill = PatternFill("solid", fgColor=PCOL[pcell.value].lstrip("#"))
                pcell.font = Font(bold=True, color="000000" if pcell.value == "P3" else "FFFFFF")
        ws = xw.book["action_tracker"]
        widths = {"A": 9, "B": 20, "C": 22, "D": 45, "E": 4, "F": 95, "G": 12, "H": 12, "I": 11, "J": 30}
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        for row in ws.iter_rows(min_row=2):
            for c in (row[3], row[5]):
                c.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"Open,In progress,Done,Won\'t do"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"H2:H{ws.max_row}")
        ws2 = xw.book["listing_summary"]
        for i, col in enumerate(ws2.iter_cols(min_row=1, max_row=min(ws2.max_row, 90)), 1):
            ws2.column_dimensions[col[0].column_letter].width = min(45, max(9, max(len(str(c.value or "")) for c in col) + 2))
        for row in ws2.iter_rows(min_row=2):
            for c in row[6:11]:
                c.number_format = '"$"#,##0'


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    pdf(OUT / "listing_actions_2026-10-06.pdf")
    xlsx(OUT / "listing_actions_2026-10-06.xlsx")
    print("ok")
