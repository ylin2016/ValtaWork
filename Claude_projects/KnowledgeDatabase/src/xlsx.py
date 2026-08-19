"""Defensive xlsx reading for the onboarding workbooks.

Several of these files are damaged in ways that break openpyxl's defaults:

* Their `<dimension>` record claims a single cell. In read_only mode openpyxl
  trusts it, so iteration silently stops after one row — no error, just missing
  data. reset_dimensions() forces a real scan.
* Some reference drawing parts (`xl/drawings/drawing1.xml`) that are absent from
  the archive, so read_only=False raises KeyError during load. That rules out
  read_only=False as the default; read_only skips drawings entirely.
* Some hold cells pointing past the end of `xl/sharedStrings.xml`
  (`Cottage 7 Onboarding Sheets.xlsx`). openpyxl indexes that table directly in
  `WorkSheetParser.parse_cell`, so the read dies with a bare `IndexError: list
  index out of range`. Both readers share that code path, so switching readers
  is not a workaround — the table itself has to tolerate the overrun.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


class _LenientStrings(list):
    """A sharedStrings table where an out-of-range reference reads as blank.

    The referenced string is genuinely absent from the archive, so no reader can
    recover it. Returning "" loses that one cell instead of the whole workbook;
    `misses` counts the losses so they can be reported rather than assumed zero.
    """

    def __init__(self, source):
        super().__init__(source)
        self.misses = 0

    def __getitem__(self, index):
        try:
            return super().__getitem__(index)
        except IndexError:
            self.misses += 1
            return ""


def open_workbook(path: str | Path):
    return load_workbook(path, data_only=True, read_only=True)


def sheet_rows(ws) -> list[list]:
    """Every row of a worksheet, immune to a bogus dimension record."""
    ws.reset_dimensions()

    # Patched per sheet, not at load: read_only builds a fresh parser around
    # ws._shared_strings on every iteration, so replacing the list is enough.
    strings = getattr(ws, "_shared_strings", None)
    if strings is not None and not isinstance(strings, _LenientStrings):
        strings = _LenientStrings(strings)
        ws._shared_strings = strings

    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    if strings is not None and strings.misses:
        wb = ws.parent
        wb._kdb_missing_strings = getattr(wb, "_kdb_missing_strings", 0) + strings.misses
        strings.misses = 0
    return rows


def missing_strings(wb) -> int:
    """Cells dropped so far because sharedStrings.xml was short."""
    return getattr(wb, "_kdb_missing_strings", 0)
